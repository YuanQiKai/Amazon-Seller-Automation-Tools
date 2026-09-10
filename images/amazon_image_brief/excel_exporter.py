from __future__ import annotations

from copy import copy
from pathlib import Path
import json

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .catalogs import CATALOG_BY_CODE
from .creative_planner import plan_description
from .models import ModuleBrief, ProductProject
from .creative_ai import recipe_for
from .final_images import final_image_path
from .languages import LANGUAGE_NAMES


NAVY = "1E2A33"
AMBER = "D88A2A"
STONE = "F4F1EA"
WARM_WHITE = "FCFBF7"
HEADER = "E8D8BF"
BORDER = "C8C2B8"
YELLOW = "FFF2CC"
GREEN = "E2F0D9"
RED = "FCE4D6"


class ExcelExporter:
    HEADERS = [
        "图片序号",
        "模块实例ID",
        "模块/图片类型",
        "PC建议尺寸（px）",
        "移动端建议尺寸（px）",
        "上传的产品参考（输入）",
        "图片生成状态 / 错误",
        "最终结果图（所选语言）",
        "用户指定Prompt",
        "设计思路和需求说明",
        "英文文案",
        "中文翻译文案",
        "德语",
        "法语",
        "意大利语",
        "西班牙语",
        "关键词映射",
        "文案版本",
        "审核状态",
        "ChatGPT自查结果/修改建议",
        "产品结构相似度",
        "OCR识别文字",
        "多语言回译检查",
        "文案溢出检查",
        "图片文字合规检查",
        "Amazon规则预检",
        "审核备注",
        "合规/备注",
        "图片版本",
        "图片审核状态",
        "图片修改意见",
        "本图主卖点",
        "卖点依据与PC/移动版式",
        "文案生成状态/失败原因",
    ]

    def export(self, project: ProductProject, briefs: list[ModuleBrief], destination: Path, poster=None) -> Path:
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        self._add_project_sheet(workbook, project)
        for channel in ("主图", "高级A+", "Brand Story", "品牌旗舰店"):
            rows = [item for item in briefs if item.channel == channel]
            if rows:
                self._add_module_sheet(workbook, project, channel, rows)
        if project.preflight_results:
            self._add_preflight_sheet(workbook, project)
        if project.copy_versions:
            self._add_version_sheet(workbook, project)
        if poster:
            self._add_poster_sheet(workbook, poster, destination.parent)
        self._add_trace_sheet(workbook, project, briefs)
        self._add_translations_sheet(workbook, project, briefs)
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
        workbook.save(destination)
        return destination

    @staticmethod
    def _language_columns(project):
        selected = set(project.active_languages())
        order = ('en', 'zh', 'de', 'fr', 'it', 'es', *LANGUAGE_NAMES)
        return list(dict.fromkeys(code for code in order if code in selected))

    def _add_translations_sheet(self, workbook, project, briefs):
        sheet = workbook.create_sheet('语言与中文对照')
        sheet.append(['模块实例', '模块名称', '语言', '当前文案', '逐行中文对照', '实际图片语言'])
        self._style_section_header(sheet, 1, 6)
        for brief in briefs:
            for code in project.active_languages():
                if code == 'zh':
                    continue
                comparison = brief.chinese_translations.get(code, '')
                if not comparison and brief.module_code != 'MAIN_WHITE':
                    comparison = '此版本暂无独立回译；通用中文文案：\n' + brief.copy_for('zh')
                sheet.append([brief.instance_id, brief.module_name, LANGUAGE_NAMES[code], brief.copy_for(code), comparison, LANGUAGE_NAMES.get(brief.image_language, brief.image_language)])
                sheet.row_dimensions[sheet.max_row].height = 100
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        for index, width in enumerate((30, 30, 18, 65, 65, 20), 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'D2'
        sheet.auto_filter.ref = sheet.dimensions

    def _add_trace_sheet(self, workbook, project, briefs):
        sheet = workbook.create_sheet('V3提示词与上下文')
        sheet.append(['实例 / 轮播帧', '字段', '内容（长内容分段）'])
        def add(iid, label, value):
            value = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, indent=2)
            for offset in range(0, max(1, len(value)), 30000):
                sheet.append([iid, label + (f' / 第{offset//30000+1}段' if len(value)>30000 else ''), value[offset:offset+30000]])
                # All AI/user content is literal text, never an Excel formula.
                for cell in sheet[sheet.max_row]:
                    cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                sheet.row_dimensions[sheet.max_row].height = 95
        add('产品', '产品理解：分析请求及回复', project.ai_context)
        for brief in briefs:
            recipe = recipe_for(project, brief.instance_id)
            key = brief.instance_id + (f' / 轮播第{brief.frame_index}/{brief.frame_count}帧' if brief.parent_id else '')
            for label, value in [('用户文案Prompt', recipe['copy_prompt']), ('用户图片Prompt', recipe['image_prompt']),
                                 ('实际文案请求Prompt', brief.effective_copy_prompt), ('实际图片请求Prompt', brief.effective_image_prompt),
                                 ('文字制作方式', recipe['render_text']), ('风格参考图与分析', {k: recipe.get(k) for k in ('style_reference_paths', 'style_analysis')}),
                                 ('PC/移动版式', brief.creative_plan)]:
                add(key, label, value)
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor=NAVY)
            cell.font = Font(color='FFFFFF', bold=True)
        for column, width in [('A', 45), ('B', 28), ('C', 115)]:
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = 'C2'
        sheet.auto_filter.ref = sheet.dimensions

    def _add_poster_sheet(self, workbook, poster, root):
        sheet = workbook.create_sheet('A+连贯海报')
        sheet.append(['A+连贯海报设计交付（PC / 移动端分别排版）'])
        sheet.append(['整体方向', poster.get('direction', '')])
        sheet.append(['重要说明', '整幅长图用于设计与审核，不能作为单张A+上传。各章是统一画布参考；多图/视频/热点/比较表/文本需按Seller Central实际模块拆分配置。'])
        sheet.append(['就绪章节', f"{poster.get('completed', 0)}/{poster.get('total', 0)}；检查各章排版状态，失败占位不应上传。"])
        sheet.append(['PC整幅', poster.get('pc', ''), poster.get('pc_size', '')])
        sheet.append(['移动端整幅', poster.get('mobile', ''), poster.get('mobile_size', '')])
        sheet.append(['序号 / 模块', '主卖点 / 状态', 'PC文件', '移动端文件', 'PC画布', '移动端画布'])
        for entry in poster.get('chapters', []):
            sheet.append([f"{entry['order']} / {entry['module']}", entry['focus']+'\n'+entry['status'], entry['pc'], entry['mobile'], entry['pc_size'], entry['mobile_size']])
            for col in (3, 4):
                cell = sheet.cell(sheet.max_row, col)
                cell.hyperlink = str(Path(cell.value).relative_to(root))
                cell.style = 'Hyperlink'
        for row in (5, 6):
            cell = sheet.cell(row, 2)
            cell.hyperlink = str(Path(cell.value).relative_to(root))
            cell.style = 'Hyperlink'
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            sheet.row_dimensions[row[0].row].height = 50
        for col, width in [('A', 35), ('B', 65), ('C', 55), ('D', 55), ('E', 20), ('F', 20)]:
            sheet.column_dimensions[col].width = width
        start = sheet.max_row + 2
        for col, device in [('A', 'pc'), ('D', 'mobile')]:
            if poster.get(device) and Path(poster[device]).is_file():
                preview = ExcelImage(poster[device])
                scale = min(450/preview.width, 1100/preview.height)
                preview.width *= scale
                preview.height *= scale
                sheet.add_image(preview, f'{col}{start}')
        sheet.freeze_panes = 'C8'

    def _add_project_sheet(self, workbook: Workbook, project: ProductProject) -> None:
        sheet = workbook.create_sheet("项目参数")
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells("A1:H1")
        sheet["A1"] = "Amazon 图片需求生成项目 V3.1"
        sheet["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 36

        fields = [
            ("项目名称", project.project_name),
            ("SKU", project.sku),
            ("ASIN", project.asin),
            ("品牌", project.brand),
            ("中文产品名", project.product_name_zh),
            ("英文产品名", project.product_name_en),
            ("品类", project.category),
            ("目标站点", project.marketplace),
            ("目标人群", project.target_audience),
            ("核心定位", project.positioning),
            ("型号", project.model),
            ("材质", project.material),
            ("基础价格", project.price),
            ("功能", project.functions),
            ("核心卖点", project.selling_points),
            ("包装清单", project.package_contents),
            ("维护/洗护", project.maintenance),
            ("保修", project.warranty),
            ("认证", project.certifications),
            ("品牌语气", project.brand_tone),
            ("视觉风格", project.visual_style),
            ("字体建议", project.font_suggestion),
            ("品牌颜色", project.brand_colors),
            ("品牌Logo", project.logo_image_path),
            ("禁用表达", project.forbidden_claims),
            ("竞品/关键词导入来源", project.imported_market_data_path),
            ("文案AI供应商", project.options.text_provider),
            ("文案协议/模型", f"{project.options.text_protocol} / {project.options.text_model}"),
            ("文案Base URL", project.options.text_base_url),
            ("图片AI供应商", project.options.image_provider),
            ("图片协议/模型", f"{project.options.image_protocol} / {project.options.image_model}"),
            ("图片Base URL", project.options.image_base_url),
            ("失败自动切换", "开启" if project.options.auto_failover else "关闭"),
            ("文案备用供应商/模型", f"{project.options.text_fallback_provider} / {project.options.text_fallback_model}"),
            ("图片备用供应商/模型", f"{project.options.image_fallback_provider} / {project.options.image_fallback_model}"),
            ("请求限流", f"{project.options.requests_per_minute} 次/分钟"),
            ("结果缓存", "开启" if project.options.cache_enabled else "关闭"),
            ("结构锁定/相似度阈值", f"{'开启' if project.options.structure_lock else '关闭'} / {project.options.similarity_threshold}"),
            ("OCR/回译/溢出/图片合规", " / ".join([
                "开" if project.options.ocr_check else "关",
                "开" if project.options.back_translation_check else "关",
                "开" if project.options.overflow_check else "关",
                "开" if project.options.image_compliance_check else "关",
            ])),
            ("Amazon规则库", project.options.rule_library_path or "内置规则库"),
            ("站点列表来源", "https://sell.amazon.com/global-selling"),
            ("模型列表来源", "https://developers.openai.com/api/docs/models"),
        ]
        sheet.append(["字段", "内容"])
        self._style_section_header(sheet, 2, 2)
        for label, value in fields:
            sheet.append([label, value])
        for label, value in project.custom_fields.items():
            sheet.append([f"自定义参数｜{label}", value])

        start = sheet.max_row + 2
        sheet.cell(start, 1, "产品变体")
        sheet.cell(start, 1).font = Font(name="Microsoft YaHei", size=13, bold=True, color=AMBER)
        sheet.append(["变体名称", "SKU", "尺寸", "颜色", "重量", "容量", "价格"])
        self._style_section_header(sheet, sheet.max_row, 7)
        for variant in project.variants:
            sheet.append([variant.name, variant.sku, variant.size, variant.color, variant.weight, variant.capacity, variant.price])

        start = sheet.max_row + 2
        sheet.cell(start, 1, "竞品信息")
        sheet.cell(start, 1).font = Font(name="Microsoft YaHei", size=13, bold=True, color=AMBER)
        sheet.append(["ASIN", "链接", "品牌", "价格", "卖点", "视觉观察", "机会点"])
        self._style_section_header(sheet, sheet.max_row, 7)
        for competitor in project.competitors:
            sheet.append([competitor.asin, competitor.url, competitor.brand, competitor.price, competitor.selling_points, competitor.visual_notes, competitor.opportunity])

        start = sheet.max_row + 2
        sheet.cell(start, 1, "关键词词库")
        sheet.cell(start, 1).font = Font(name="Microsoft YaHei", size=13, bold=True, color=AMBER)
        sheet.append(["关键词", "语言", "意图", "优先级", "备注"])
        self._style_section_header(sheet, sheet.max_row, 5)
        for keyword in project.keywords:
            sheet.append([keyword.term, keyword.language, keyword.intent, keyword.priority, keyword.notes])

        start = sheet.max_row + 2
        sheet.cell(start, 1, "已确认模块顺序与Prompt")
        sheet.cell(start, 1).font = Font(name="Microsoft YaHei", size=13, bold=True, color=AMBER)
        sheet.append(["顺序", "渠道", "模块", "模块代码", "模块实例ID", "用户指定Prompt"])
        self._style_section_header(sheet, sheet.max_row, 6)
        for index, instance in enumerate(project.normalized_module_instances(), 1):
            spec = CATALOG_BY_CODE.get(instance.module_code)
            if spec:
                sheet.append(
                    [
                        index,
                        spec.channel,
                        spec.name,
                        instance.module_code,
                        instance.instance_id,
                        instance.custom_prompt or project.module_prompts.get(instance.instance_id, ""),
                    ]
                )

        sheet.append(['品牌 Slogan', project.brand_slogan])
        sheet.append(['目标文案语言', '、'.join(LANGUAGE_NAMES[code] for code in project.active_languages())])
        sheet.append(['图片文案语言', LANGUAGE_NAMES[project.render_language()]])
        sheet.append(['品牌 Logo', project.logo_image_path])
        self._place_image(sheet, project.logo_image_path, sheet.max_row, 3)
        sheet.row_dimensions[sheet.max_row].height = 110
        thin = Side(style="thin", color=BORDER)
        for row in sheet.iter_rows():
            for cell in row:
                updated_font = copy(cell.font)
                updated_font.name = "Microsoft YaHei"
                updated_font.sz = updated_font.sz or 10
                cell.font = updated_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.value is not None:
                    cell.border = Border(bottom=thin)
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 64
        for column in range(3, 9):
            sheet.column_dimensions[get_column_letter(column)].width = 30
        sheet.freeze_panes = "A3"

    @staticmethod
    def _style_section_header(sheet, row: int, columns: int) -> None:
        for column in range(1, columns + 1):
            cell = sheet.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=HEADER)
            cell.font = Font(name="Microsoft YaHei", bold=True, color=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _add_module_sheet(self, workbook: Workbook, project: ProductProject, channel: str, rows: list[ModuleBrief]) -> None:
        languages = self._language_columns(project)
        headers = self.HEADERS[:10] + [LANGUAGE_NAMES[code] + '文案' for code in languages] + self.HEADERS[16:]
        offset = len(languages) - 6
        sheet = workbook.create_sheet(channel[:31])
        sheet.sheet_view.showGridLines = False
        end_column = get_column_letter(len(headers))
        sheet.merge_cells(f"A1:{end_column}1")
        sheet["A1"] = f"{project.project_name}｜{channel}图片需求"
        sheet["A1"].font = Font(name="Microsoft YaHei", size=17, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 34
        sheet.merge_cells(f"A2:{end_column}2")
        sheet["A2"] = (
            "尺寸为建议制作尺寸，上传前以目标站点 Seller Central 实时提示为准。只展示最终结果图，不生成草稿/失败占位图；"
            "产品结构、颜色、文字、翻译、声明与审核状态必须人工终审。"
        )
        sheet["A2"].fill = PatternFill("solid", fgColor=YELLOW)
        sheet["A2"].font = Font(name="Microsoft YaHei", size=10, bold=True, color="7A4A00")
        sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[2].height = 32

        for column, header in enumerate(headers, 1):
            cell = sheet.cell(3, column, header)
            cell.fill = PatternFill("solid", fgColor=HEADER)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[3].height = 45

        for row_index, brief in enumerate(rows, 4):
            values = [
                brief.sequence,
                brief.instance_id,
                brief.module_name,
                brief.pc_size,
                brief.mobile_size,
                "",
                brief.image_generation_status + ('：' + brief.image_generation_error if brief.image_generation_error else ''),
                "",
                recipe_for(project, brief.instance_id)['image_prompt'],
                brief.design_brief,
                *(brief.copy_for(code) for code in languages),
                " / ".join(brief.keywords),
                f"V{brief.copy_version}",
                brief.review_status,
                brief.self_check_result,
                "" if brief.image_similarity is None else f"{brief.image_similarity:.1f}/100",
                brief.ocr_text,
                brief.back_translation_result,
                brief.overflow_result,
                brief.image_qa_result,
                brief.rule_preflight_result,
                brief.review_notes,
                brief.compliance_note,
                f"V{brief.image_version}",
                brief.image_review_status,
                brief.image_revision_notes,
                brief.creative_plan.get('focus', ''),
                plan_description(brief.creative_plan) if brief.creative_plan else '',
                brief.generation_status + ('：' + brief.generation_error if brief.generation_error else ''),
            ]
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row_index, column, value)
                cell.fill = PatternFill("solid", fgColor=WARM_WHITE if row_index % 2 == 0 else STONE)
                cell.font = Font(name="Arial" if 11 <= column <= 16 else "Microsoft YaHei", size=9, color=NAVY)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            status_fill = {"已通过": GREEN, "需修改": RED}.get(brief.review_status, YELLOW)
            sheet.cell(row_index, 19 + offset).fill = PatternFill("solid", fgColor=status_fill)
            image_status_fill = {"满意": GREEN, "不满意-待修改": RED, "待复核": YELLOW}.get(brief.image_review_status, YELLOW)
            sheet.cell(row_index, 30 + offset).fill = PatternFill("solid", fgColor=image_status_fill)
            sheet.cell(row_index, 1).font = Font(name="Microsoft YaHei", size=11, bold=True, color=AMBER)
            sheet.row_dimensions[row_index].height = 125
            self._place_image(sheet, brief.reference_image, row_index, 6)
            result_image = final_image_path(brief)
            self._place_image(sheet, str(result_image) if result_image else '', row_index, 8)

        thin = Side(style="thin", color=BORDER)
        for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = Border(bottom=thin, right=thin)
        widths = [15, 30, 25, 18, 22, 25, 25, 25, 38, 46, 30, 30, 30, 30, 30, 30, 25, 12, 14, 42, 18, 32, 48, 38, 42, 38, 32, 34, 12, 18, 42, 38, 68, 40]
        widths = widths[:10] + [34] * len(languages) + widths[16:]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:{end_column}{sheet.max_row}"

    def _add_preflight_sheet(self, workbook: Workbook, project: ProductProject) -> None:
        sheet = workbook.create_sheet("Amazon规则预检")
        sheet.sheet_view.showGridLines = False
        report = project.preflight_results if isinstance(project.preflight_results, dict) else {}
        sheet.merge_cells("A1:F1")
        sheet["A1"] = (
            f"Amazon 导出前预检｜规则库 {report.get('library_version', '-')}｜"
            f"站点 {report.get('marketplace_code', '-')}｜品类 {report.get('category_key', '-')}"
        )
        sheet["A1"].font = Font(name="Microsoft YaHei", size=15, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 34
        headers = ["级别", "规则ID", "作用范围", "检查结果", "处理建议", "依据"]
        sheet.append(headers)
        self._style_section_header(sheet, 2, len(headers))
        for item in report.get("findings", []):
            sheet.append([
                item.get("severity", ""), item.get("rule_id", ""), item.get("scope", ""),
                item.get("message", ""), item.get("suggestion", ""), item.get("source_url", ""),
            ])
            color = {"错误": RED, "警告": YELLOW, "信息": GREEN}.get(item.get("severity"), WARM_WHITE)
            sheet.cell(sheet.max_row, 1).fill = PatternFill("solid", fgColor=color)
        for row in sheet.iter_rows(min_row=3):
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=9, color=NAVY)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=Side(style="thin", color=BORDER))
        for index, width in enumerate([12, 20, 22, 48, 48, 52], 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A3"
        sheet.auto_filter.ref = f"A2:F{max(2, sheet.max_row)}"

    def _add_version_sheet(self, workbook: Workbook, project: ProductProject) -> None:
        languages = self._language_columns(project)
        sheet = workbook.create_sheet("文案版本记录")
        sheet.sheet_view.showGridLines = False
        headers = ["模块实例ID", "模块代码", "模块名称", "版本", "记录时间", "变更原因", *(LANGUAGE_NAMES[code] for code in languages)]
        sheet.append(headers)
        self._style_section_header(sheet, 1, len(headers))
        for instance in project.normalized_module_instances():
            spec = CATALOG_BY_CODE.get(instance.module_code)
            for version in project.copy_versions.get(instance.instance_id, project.copy_versions.get(instance.module_code, [])):
                copy_text = version.get("copy", {})
                sheet.append([
                    instance.instance_id,
                    instance.module_code,
                    spec.name if spec else instance.module_code,
                    f"V{version.get('version', 1)}",
                    version.get("time", ""),
                    version.get("reason", ""),
                    *(copy_text.get(code, '') for code in languages),
                ])
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=9, color=NAVY)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=Side(style="thin", color=BORDER))
        widths = [30, 24, 26, 10, 20, 28, 34, 34, 34, 34, 34, 34]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"

    @staticmethod
    def _place_image(sheet, image_path: str, row: int, column: int) -> None:
        path = Path(image_path) if image_path else None
        if not path or not path.is_file():
            return
        image = ExcelImage(str(path))
        image.width = 165
        image.height = 108
        sheet.add_image(image, f"{get_column_letter(column)}{row}")
