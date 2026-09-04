from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Competitor, Keyword


NAVY = "1E2A33"
AMBER = "D88A2A"
PALE = "F4F1EA"
INPUT = "FFF2CC"


COMPETITOR_HEADERS = ["ASIN", "链接", "品牌", "价格", "卖点", "视觉观察", "可突破机会点"]
KEYWORD_HEADERS = ["关键词", "语言", "搜索意图", "优先级", "备注"]


def _normalized(value: object) -> str:
    return str(value or "").strip().replace(" ", "").lower()


def _header_indexes(row: Iterable[object], aliases: dict[str, set[str]]) -> dict[str, int]:
    result: dict[str, int] = {}
    normalized = [_normalized(value) for value in row]
    for field, names in aliases.items():
        accepted = {_normalized(name) for name in names}
        for index, value in enumerate(normalized):
            if value in accepted:
                result[field] = index
                break
    return result


class MarketWorkbookManager:
    COMP_ALIASES = {
        "asin": {"ASIN", "竞品ASIN"},
        "url": {"链接", "URL", "竞品链接"},
        "brand": {"品牌", "Brand"},
        "price": {"价格", "Price"},
        "selling_points": {"卖点", "核心卖点", "Selling Points"},
        "visual_notes": {"视觉观察", "图片观察", "Visual Notes"},
        "opportunity": {"可突破机会点", "机会点", "Opportunity"},
    }
    KEYWORD_ALIASES = {
        "term": {"关键词", "Keyword", "Search Term"},
        "language": {"语言", "Language"},
        "intent": {"搜索意图", "意图", "Intent"},
        "priority": {"优先级", "Priority"},
        "notes": {"备注", "Notes"},
    }

    def export_template(self, destination: Path) -> Path:
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        guide = workbook.active
        guide.title = "使用说明"
        guide.sheet_view.showGridLines = False
        guide.append(["竞品与关键词导入模板"])
        guide.append(["填写方式", "不要修改“竞品信息”和“关键词词库”的表名及第1行字段名；从第2行开始填写。"])
        guide.append(["竞品", "ASIN或链接至少填写一项。卖点、视觉观察和机会点可使用换行。"])
        guide.append(["关键词", "优先级填写1-5，5为最高；语言建议使用en/zh/de/fr/it/es。"])
        guide.append(["导入", "保存为.xlsx后，在程序“竞品与关键词”页选择Excel导入。"])
        guide.column_dimensions["A"].width = 18
        guide.column_dimensions["B"].width = 90
        guide["A1"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
        for row in guide.iter_rows(min_row=2, max_row=5, min_col=1, max_col=2):
            row[0].font = Font(name="Arial", bold=True, color=NAVY)
            row[1].alignment = Alignment(wrap_text=True, vertical="top")

        competitor = workbook.create_sheet("竞品信息")
        keyword = workbook.create_sheet("关键词词库")
        example = workbook.create_sheet("填写示例")
        for sheet, headers in ((competitor, COMPETITOR_HEADERS), (keyword, KEYWORD_HEADERS)):
            sheet.sheet_view.showGridLines = False
            sheet.append(headers)
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in range(2, 202):
                for column in range(1, len(headers) + 1):
                    sheet.cell(row, column).fill = PatternFill("solid", fgColor=INPUT)
                    sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:{chr(64 + len(headers))}201"
        for column, width in enumerate([18, 46, 20, 16, 42, 42, 42], 1):
            competitor.column_dimensions[chr(64 + column)].width = width
        for column, width in enumerate([32, 14, 20, 12, 42], 1):
            keyword.column_dimensions[chr(64 + column)].width = width

        language_validation = DataValidation(type="list", formula1='"en,zh,de,fr,it,es"', allow_blank=True)
        priority_validation = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
        keyword.add_data_validation(language_validation)
        keyword.add_data_validation(priority_validation)
        language_validation.add("B2:B201")
        priority_validation.add("D2:D201")

        example.append(COMPETITOR_HEADERS)
        example.append(["B0EXAMPLE", "https://www.amazon.de/dp/B0EXAMPLE", "示例品牌", "€89.99", "顺滑万向轮", "尺寸图清楚，但收纳结构表达较弱", "突出隔层与扩容前后对比"])
        example.append([])
        example.append(KEYWORD_HEADERS)
        example.append(["Koffer mit Rollen", "de", "品类词", 5, "自然用于尺寸或场景图"])
        example.sheet_view.showGridLines = False
        for row_number in (1, 4):
            for cell in example[row_number]:
                cell.fill = PatternFill("solid", fgColor=AMBER)
                cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in example.iter_rows(min_row=2, max_row=5, min_col=1, max_col=7):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        example.row_dimensions[2].height = 44
        example.row_dimensions[5].height = 36
        for column, width in enumerate([18, 46, 20, 16, 34, 38, 38], 1):
            example.column_dimensions[chr(64 + column)].width = width
        workbook.save(destination)
        workbook.close()
        return destination

    def import_data(self, source: Path) -> tuple[list[Competitor], list[Keyword]]:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            competitor_sheet = self._find_sheet(workbook, ("竞品信息", "竞品", "competitors"))
            keyword_sheet = self._find_sheet(workbook, ("关键词词库", "关键词", "keywords"))
            competitors = self._read_competitors(competitor_sheet) if competitor_sheet else []
            keywords = self._read_keywords(keyword_sheet) if keyword_sheet else []
        finally:
            workbook.close()
        if not competitors and not keywords:
            raise ValueError("没有识别到可导入数据。请使用程序导出的模板，并从第2行开始填写。")
        return competitors, keywords

    @staticmethod
    def _find_sheet(workbook, accepted_names: tuple[str, ...]):
        normalized = {_normalized(name) for name in accepted_names}
        return next((sheet for sheet in workbook.worksheets if _normalized(sheet.title) in normalized), None)

    def _read_competitors(self, sheet) -> list[Competitor]:
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, ())
        indexes = _header_indexes(headers, self.COMP_ALIASES)
        result: list[Competitor] = []
        for row in rows:
            values = {field: str(row[index] or "").strip() if index < len(row) else "" for field, index in indexes.items()}
            if not values.get("asin") and not values.get("url"):
                continue
            result.append(Competitor(**values))
        return result

    def _read_keywords(self, sheet) -> list[Keyword]:
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, ())
        indexes = _header_indexes(headers, self.KEYWORD_ALIASES)
        result: list[Keyword] = []
        for row in rows:
            values = {field: row[index] if index < len(row) else "" for field, index in indexes.items()}
            term = str(values.get("term") or "").strip()
            if not term:
                continue
            try:
                priority = max(1, min(5, int(values.get("priority") or 3)))
            except (TypeError, ValueError):
                priority = 3
            result.append(
                Keyword(
                    term=term,
                    language=str(values.get("language") or "en").strip(),
                    intent=str(values.get("intent") or "feature").strip(),
                    priority=priority,
                    notes=str(values.get("notes") or "").strip(),
                )
            )
        return result
