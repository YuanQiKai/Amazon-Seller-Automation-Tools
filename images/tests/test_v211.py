from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path
import tempfile
import threading
import time
import tkinter as tk
import unittest
from unittest.mock import patch
import zipfile

from PIL import Image, ImageDraw
from openpyxl import load_workbook

from amazon_image_brief.aplus_poster import build_poster
from amazon_image_brief.brief_generator import BriefGenerator
from amazon_image_brief.creative_planner import plan_modules, layout_prompt, apply_layout
from amazon_image_brief.generation_control import GenerationControl, GenerationCancelled
from amazon_image_brief.ai_runtime import RequestRateLimiter
from amazon_image_brief.gui import AmazonImageBriefApp
from amazon_image_brief.image_composer import create_german_composite
from amazon_image_brief.layout_editor import LayoutEditor
from amazon_image_brief.models import ModuleInstance, ProductProject, LANGUAGES
from amazon_image_brief.service import GenerationService
from amazon_image_brief.profile_store import ProfileStore
from amazon_image_brief.typography import PRESETS, typeset, validate_element
from test_v28 import PlannedAI, project_at
from test_v210 import ImageAI


def await_state(control, state, timeout=8):
    until = time.monotonic()+timeout
    while time.monotonic() < until:
        if control.state == state:
            return
        time.sleep(.02)
    raise AssertionError(f'Expected {state}, got {control.state}')


def poster_project(root):
    project = project_at(root/'out')
    project.options.aplus_continuous = True
    project.options.aplus_direction = 'Warm neutral editorial story; keep typography consistent.'
    project.options.generate_german_composites = True
    project.module_instances += [ModuleInstance('a1', 'APLUS_FULL_IMAGE', '高级A+'), ModuleInstance('a2', 'APLUS_FULL_IMAGE', '高级A+'),
                                 ModuleInstance('brand', 'BS_BRAND_FOCUS', 'Brand Story'), ModuleInstance('store', 'STORE_HERO', '品牌旗舰店')]
    reference = root/'reference.jpg'
    Image.new('RGB', (1200, 1200), '#f5f3ed').save(reference)
    project.product_image_paths = [str(reference)]
    return project


def concise(briefs):
    for brief in briefs:
        if brief.module_code != 'MAIN_WHITE':
            values = ['Reisen mit Komfort', 'Details für unterwegs', 'Leichtes Gepäck', 'Praktische Aufteilung', 'Sanftes Rollen', 'Klare Ordnung']
            brief.copy = {lang: '\n'.join(values[:len(brief.creative_plan['copy_slots'])]) for lang in LANGUAGES}
    return briefs


class ControlTests(unittest.TestCase):
    def test_copy_pause_keeps_inflight_result_and_continues_without_duplicate(self):
        project = project_at('unused')
        project.options.optimize_copy_with_ai = True
        events, result = [], []
        control = GenerationControl('copy')
        control.start()
        class Client(PlannedAI):
            def generate_json(self, prompt, model):
                payload = super().generate_json(prompt, model)
                if len(self.requests) == 1:
                    control.pause()
                return payload
        client = Client()
        generator = BriefGenerator(client, events.append, control.checkpoint)
        thread = threading.Thread(target=lambda: result.extend(generator.generate(project)), daemon=True)
        thread.start()
        try:
            await_state(control, 'paused')
            self.assertEqual(1, len(client.requests))
            self.assertTrue(any(e['instance_id'] == 'detail1' and e['finished'] for e in events))
            self.assertEqual([], result)
        finally:
            control.resume()
            thread.join(10)
        self.assertFalse(thread.is_alive())
        self.assertEqual(2, len(client.requests))
        self.assertEqual(3, len(result))

    def test_image_pause_keeps_checkpoint_then_resumes_remaining(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            project = project_at(root/'out')
            project.options.generate_ai_images = True
            source = root/'source.jpg'
            Image.new('RGB', (1200, 1200), 'white').save(source)
            project.product_image_paths = [str(source)]
            control = GenerationControl('image')
            control.start()
            client = ImageAI(before_call=lambda index: control.pause() if index == 0 else None)
            events, results = [], []
            service = GenerationService(root, client, on_image_progress=events.append, image_control=control.checkpoint)
            thread = threading.Thread(target=lambda: results.append(service.run(project)), daemon=True)
            thread.start()
            try:
                await_state(control, 'paused')
                self.assertEqual(1, len(client.image_calls))
                self.assertEqual('complete', events[-1]['stage'])
                self.assertTrue(Path(events[-1]['output_dir'], 'partial_briefs.json').is_file())
            finally:
                control.resume()
                # The resume assertion includes JPEG QA and Excel packaging on
                # disk, not just the control signal; allow slower Windows disks.
                thread.join(60)
            self.assertFalse(thread.is_alive())
            self.assertEqual(3, len(client.image_calls))
            self.assertTrue(Path(results[0]['excel']).is_file())

    def test_controls_independent_and_close_unblocks_pause(self):
        copy, image = GenerationControl('copy'), GenerationControl('image')
        copy.start()
        image.start()
        copy.pause()
        image.checkpoint()
        errors = []
        def wait():
            try:
                copy.checkpoint()
            except GenerationCancelled as exc:
                errors.append(str(exc))
        thread = threading.Thread(target=wait)
        thread.start()
        await_state(copy, 'paused')
        copy.close()
        thread.join(3)
        self.assertEqual('running', image.state)
        self.assertEqual(1, len(errors))

    def test_rate_limit_wait_observes_pause(self):
        limiter = RequestRateLimiter(1)
        limiter.acquire()
        control = GenerationControl('image')
        control.start()
        control.pause()
        errors = []
        def acquire():
            try:
                limiter.acquire(control.checkpoint)
            except GenerationCancelled:
                errors.append(True)
        thread = threading.Thread(target=acquire)
        thread.start()
        await_state(control, 'paused')
        control.close()
        thread.join(3)
        self.assertEqual([True], errors)
        self.assertEqual(1, len(limiter._timestamps))


class LayoutAndPosterTests(unittest.TestCase):
    def test_smart_layout_elements_and_manual_coordinates_roundtrip(self):
        project = project_at('unused')
        plans = plan_modules(project)
        self.assertEqual('hotspots', plans['detail1']['layout'])
        self.assertEqual([], plans['white']['copy_slots'])
        for device in ('pc', 'mobile'):
            elements = plans['detail1'][device]['elements']
            self.assertEqual(4, len(elements))
            for item in elements:
                validate_element(item)
            self.assertGreater(elements[0]['font_ratio'], elements[1]['font_ratio'])
        plan = plans['detail1']
        plan['manual'] = True
        plan['pc']['elements'][0]['box'] = [.05, .05, .90, .15]
        project.creative_plans = plans
        restored = ProductProject.from_dict(project.to_dict())
        self.assertEqual([.05, .05, .90, .15], plan_modules(restored)['detail1']['pc']['elements'][0]['box'])

    def test_each_preset_has_valid_separate_pc_mobile_slots(self):
        for preset in PRESETS:
            plan = {'copy_slots': ['标题', '副标题', '功能1', '功能2', '功能3', '功能4']}
            apply_layout(plan, preset)
            for device in ('pc', 'mobile'):
                for element in plan[device]['elements']:
                    validate_element(element)
            self.assertNotEqual(plan['pc']['elements'], plan['mobile']['elements'])

    def test_typesetting_checks_overflow_and_does_not_touch_white_image(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            src, out = root/'source.jpg', root/'out.jpg'
            Image.new('RGB', (1200, 1200), 'white').save(src)
            plan = plan_modules(project_at(root))['detail1']
            create_german_composite(src, out, 'Titel\nUntertitel\nDetail eins\nDetail zwei', creative_plan=plan)
            self.assertNotEqual(src.read_bytes(), out.read_bytes())
            with self.assertRaisesRegex(ValueError, '溢出'):
                typeset((1200, 1200), 'W'*4000, plan, validate_only=True)
            create_german_composite(src, out, 'ignored', skip_text=True, creative_plan=plan)
            self.assertEqual(src.read_bytes(), out.read_bytes())

    def test_poster_only_aplus_and_uses_whole_story_in_every_prompt(self):
        with tempfile.TemporaryDirectory() as temp:
            project = poster_project(Path(temp))
            plans = plan_modules(project)
            for iid, plan in plans.items():
                self.assertEqual(iid in ('a1', 'a2'), bool(plan.get('poster')))
            self.assertEqual(['a1', 'a2'], [chapter['id'] for chapter in plans['a1']['poster']['story']])
            self.assertIn('Warm neutral', layout_prompt(plans['a2']))
            self.assertIn('CONTINUOUS A+', layout_prompt(plans['a1']))
            project.options.aplus_continuous = False
            project.creative_plans = plans
            self.assertFalse(any(p.get('poster') for p in plan_modules(project).values()))

    def test_poster_export_dimensions_excel_zip_and_local_recompose(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            project = poster_project(root)
            briefs = concise(BriefGenerator().generate(project))
            for brief in briefs:
                brief.ai_effect_image = project.product_image_paths[0]
            client = ImageAI()
            service = GenerationService(root, client)
            result = service.refresh_output(project, briefs, root/'result')
            manifest = result['aplus_poster']
            self.assertEqual(2, manifest['completed'])
            self.assertEqual([], manifest['warnings'])
            with Image.open(manifest['pc']) as image:
                self.assertEqual((1464, 1200), image.size)
            with Image.open(manifest['mobile']) as image:
                self.assertEqual((1200, 1800), image.size)
            with zipfile.ZipFile(result['image_zip']) as package:
                self.assertTrue(any('aplus-pc-full.jpg' in n for n in package.namelist()))
                self.assertTrue(any('mobile-final.jpg' in n for n in package.namelist()))
            workbook = load_workbook(result['excel'])
            self.assertIn('A+连贯海报', workbook.sheetnames)
            self.assertEqual('1464×600', workbook['A+连贯海报']['E8'].value)
            workbook.close()
            a1 = next(b for b in briefs if b.instance_id == 'a1')
            self.assertEqual(manifest['chapters'][0]['pc'], a1.german_composite_image)
            old_version = manifest['version']
            a1.copy['de'] = 'Neue Reise\nMit Komfort\nLeicht unterwegs\nKlare Ordnung'
            regenerated = service.refresh_output(project, briefs, root/'result')['aplus_poster']
            self.assertNotEqual(old_version, regenerated['version'])
            self.assertTrue(Path(manifest['pc']).is_file())
            self.assertEqual([], client.image_calls)
            self.assertEqual([], client.requests)

    def test_group_regeneration_preserves_other_channels_and_real_product_reference(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            project = poster_project(root)
            briefs = concise(BriefGenerator().generate(project))
            for brief in briefs:
                brief.ai_effect_image = project.product_image_paths[0]
            others = {b.instance_id: (b.ai_effect_image, b.image_version, b.copy) for b in briefs if b.channel != '高级A+'}
            client = ImageAI()
            service = GenerationService(root, client)
            refs = []
            original = client.generate_image
            def capture(**kw):
                refs.append(kw.get('reference_images', []))
                return original(**kw)
            with patch.object(client, 'generate_image', side_effect=capture):
                result = service.regenerate_images(project, briefs, ['a1', 'a2'], root/'out', {'a1': 'warm', 'a2': 'warm'})
            self.assertEqual(2, len(client.image_calls))
            self.assertEqual(project.product_image_paths, refs[1])
            self.assertEqual([], result['regeneration_failures'])
            for b in briefs:
                if b.channel != '高级A+':
                    self.assertEqual(others[b.instance_id], (b.ai_effect_image, b.image_version, b.copy))

    def test_partial_poster_is_explicitly_marked_not_success(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            project = poster_project(root)
            briefs = concise(BriefGenerator().generate(project))
            manifest = build_poster(project, briefs, root/'out')
            self.assertEqual(0, manifest['completed'])
            self.assertTrue(all('占位' in entry['status'] for entry in manifest['chapters']))

    def test_poster_assets_are_copied_into_persistent_profiles(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            project = poster_project(root)
            path = project.product_image_paths[0]
            store = ProfileStore(root/'profiles')
            store.save({'project': project.to_dict(), 'output': {'result': {'aplus_poster':
                       {'pc': path, 'mobile': path, 'chapters': [{'pc': path, 'mobile': path}], 'files': [path]}}}})
            restored = store.load()['output']['result']['aplus_poster']
            self.assertTrue(Path(restored['pc']).is_relative_to(root/'profiles'/'assets'))
            self.assertTrue(Path(restored['chapters'][0]['mobile']).is_file())


class DesktopTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = tk.Tk()
        self.root.withdraw()
        self.app = AmazonImageBriefApp(self.root, Path(self.temp.name))

    def tearDown(self):
        for control in self.app.generation_controls.values():
            control.close()
        for callback in self.root.tk.call('after', 'info'):
            self.root.after_cancel(callback)
        self.root.destroy()
        self.temp.cleanup()

    def test_controls_and_saved_aplus_settings(self):
        app = self.app
        app._start_controls('draft')
        event = app.events.get_nowait()[1]
        app._control_state(event)
        self.assertEqual('normal', str(app.control_widgets['copy'][0]['state']))
        self.assertEqual('disabled', str(app.control_widgets['image'][0]['state']))
        app.generation_controls['copy'].pause()
        app._control_state(app.events.get_nowait()[1])
        self.assertEqual('normal', str(app.control_widgets['copy'][1]['state']))
        app.generation_controls['copy'].resume()
        app._finish_controls()
        app.aplus_continuous_var.set(True)
        app.aplus_direction_var.set('统一黑金海报')
        snapshot = app.profiles.capture()
        app.aplus_continuous_var.set(False)
        app.profiles.restore(snapshot)
        self.assertTrue(app.aplus_continuous_var.get())
        self.assertEqual('统一黑金海报', app.aplus_direction_var.get())

    def test_element_editor_saves_positions_and_rejects_invalid_bounds(self):
        plan = plan_modules(project_at('unused'))['detail1']
        saved = []
        editor = LayoutEditor(self.root, plan, saved.append)
        editor.vars['x'].set('90')
        self.assertFalse(editor.update_element())
        editor.vars['x'].set('5')
        editor.vars['w'].set('90')
        self.assertTrue(editor.update_element())
        editor.apply()
        self.assertTrue(saved[0]['manual'])
        self.assertEqual(.05, saved[0]['pc']['elements'][0]['box'][0])
        self.assertNotEqual(saved[0]['pc']['elements'][0]['box'], plan['pc']['elements'][0]['box'])


if __name__ == '__main__':
    unittest.main()
