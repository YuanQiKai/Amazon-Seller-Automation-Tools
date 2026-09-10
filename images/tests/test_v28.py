from __future__ import annotations

from copy import deepcopy
from contextlib import redirect_stdout
from io import StringIO, BytesIO
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
import tkinter as tk
from tkinter import ttk

from PIL import Image, ImageChops
from openpyxl import load_workbook
from amazon_image_brief.asset_gallery import append_image_paths
from amazon_image_brief.api_debug import APIDebugLogger
from amazon_image_brief.ai_client import OpenAIClient, OpenAIError
from amazon_image_brief.brief_generator import BriefGenerator
from amazon_image_brief.creative_planner import plan_modules, LAYOUTS
from amazon_image_brief.gui import AmazonImageBriefApp
from amazon_image_brief.image_composer import create_german_composite
from amazon_image_brief.models import ProductProject, ModuleInstance, GenerationOptions, LANGUAGES
from amazon_image_brief.rules import RuleLibrary
from amazon_image_brief.service import GenerationService


def project_at(root):
    return ProductProject(project_name='V28 test', brand='Example', product_name_zh='旅行箱', category='Luggage',
                          copy_languages=['en', 'de', 'fr', 'it', 'es'],
                          selling_points='轻量箱体\n双排万向轮\n分区收纳\n可调拉杆', functions='平稳推行\n便于分类整理',
                          module_instances=[ModuleInstance('white', 'MAIN_WHITE', '主图'), ModuleInstance('detail1', 'MAIN_DETAIL', '主图'), ModuleInstance('detail2', 'MAIN_DETAIL', '主图')],
                          options=GenerationOptions(context_before_generation=False, ai_plan_before_copy=False, optimize_copy_with_ai=False, generate_ai_images=False, generate_german_composites=False, api_debug_enabled=False, output_root=str(root)))


class PlannedAI:
    available = True
    def __init__(self, fail_first=False):
        self.requests = []
        self.debug_context = {}
        self.fail_first = fail_first

    def generate_json(self, prompt, model):
        self.requests.append(prompt)
        module = json.loads(prompt.split('\nModule: ')[-1])
        iid = module['instance_id']
        if self.fail_first and iid == 'detail1':
            raise OpenAIError('模拟 HTTP 429')
        slots = module['creative_plan']['copy_slots']
        languages = module.get('requested_languages', LANGUAGES)
        return {'instance_id': iid, 'copy': {lang: '\n'.join(f'{iid} {lang} {slot}' for slot in slots) for lang in languages},
                'chinese_translations': {lang: '\n'.join(f'{iid} {lang}中文 {slot}' for slot in slots) for lang in languages if lang != 'zh'},
                'design_brief_zh': '按卖点安排局部特写与文字留白', 'image_prompt_en': 'Product close-up with reserved copy region, no text.'}


class WorkflowTests(unittest.TestCase):
    def test_append_uploads_keep_previous_and_deduplicate(self):
        with tempfile.TemporaryDirectory() as temp:
            a, b, c = [str(Path(temp) / name) for name in ('a.png', 'b.png', 'c.png')]
            self.assertEqual([a, b, c], append_image_paths([a, b], [b, c, a]))

    def test_plans_differentiate_repeated_modules_and_manual_plan_roundtrip(self):
        project = project_at('outputs')
        plans = plan_modules(project)
        self.assertEqual([], plans['white']['copy_slots'])
        self.assertNotEqual(plans['detail1']['angle'], plans['detail2']['angle'])
        self.assertNotEqual(plans['detail1']['focus'], plans['detail2']['focus'])
        plans['detail1'].update(manual=True, focus='指定轮组近景', layout='right')
        project.creative_plans = plans
        restored = ProductProject.from_dict(project.to_dict())
        self.assertEqual('指定轮组近景', plan_modules(restored)['detail1']['focus'])

    def test_each_module_has_progress_and_fresh_generation_beats_old_overrides(self):
        project = project_at('outputs')
        project.options.optimize_copy_with_ai = True
        project.copy_overrides['detail1'] = {lang: 'OLD' for lang in LANGUAGES}
        client = PlannedAI()
        events = []
        briefs = BriefGenerator(client, events.append).generate(project)
        self.assertEqual(2, len(client.requests))
        self.assertEqual('首图无文案', briefs[0].generation_status)
        self.assertIn('AI完成', briefs[1].generation_status)
        self.assertNotIn('OLD', briefs[1].copy['de'])
        self.assertEqual(4, len(briefs[1].copy['de'].splitlines()))
        self.assertEqual(3, events[-1]['done'])
        self.assertTrue(any(event['status'] == '正在生成' for event in events))

    def test_failure_visible_and_other_modules_continue(self):
        project = project_at('outputs')
        project.options.optimize_copy_with_ai = True
        briefs = BriefGenerator(PlannedAI(True)).generate(project)
        self.assertIn('429', briefs[1].generation_error)
        self.assertIn('生成失败', briefs[1].generation_status)
        self.assertIn('AI完成', briefs[2].generation_status)

    def test_duplicate_headline_is_rejected(self):
        project = project_at('outputs')
        briefs = BriefGenerator().generate(project)
        first, second = briefs[1:]
        first.copy = {lang: 'same headline\nsupport\ndetail A\ndetail B' for lang in LANGUAGES}
        response = {'copy': dict(first.copy), 'design_brief_zh': '布局', 'image_prompt_en': 'no text'}
        with self.assertRaisesRegex(OpenAIError, '重复'):
            BriefGenerator._validate_copy_result(second, response, [first])

    def test_low_resolution_reference_does_not_block_export_and_output_errors_remain(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            reference = root / '300px.png'
            Image.new('RGB', (300, 300), 'white').save(reference)
            project = project_at(root / 'out')
            project.product_image_paths = [str(reference)]
            briefs = BriefGenerator().generate(project)
            report = RuleLibrary.load().preflight(project, briefs)
            self.assertFalse(report.errors)
            self.assertTrue(any(item.rule_id == 'REF-SIZE-001' for item in report.findings))
            output = GenerationService(root).run(project, briefs=briefs)
            self.assertTrue(Path(output['excel']).is_file())
            with open(output['excel'], 'rb') as handle:
                workbook = load_workbook(handle)
                self.assertEqual(34, workbook['主图'].max_column)
                self.assertEqual('本图主卖点', workbook['主图'].cell(3, 32).value)
                self.assertEqual(briefs[1].creative_plan['focus'], workbook['主图'].cell(5, 32).value)
                workbook.close()
            briefs[0].ai_effect_image = str(reference)
            self.assertTrue(RuleLibrary.load().preflight(project, briefs, check_outputs=True).errors)
            briefs[0].copy['de'] = 'Not allowed'
            self.assertTrue(RuleLibrary.load().preflight(project, briefs).errors)

    def test_http_error_reaches_ui_json_with_context_and_redaction(self):
        events = []
        options = GenerationOptions(api_debug_enabled=False, max_retries=0)
        client = OpenAIClient(text_api_key='real-test-key', options=options, on_api_event=events.append)
        client.debug_context.update(instance_id='detail1')
        error = HTTPError('https://example.test/v1/models', 403, 'Forbidden', {'CF-Ray': 'test-ray'}, BytesIO(b'{"error":"1010"}'))
        with patch('urllib.request.urlopen', side_effect=error):
            with self.assertRaises(OpenAIError):
                client.list_models('text')
        error.close()
        self.assertEqual(['request', 'response'], [event['phase'] for event in events])
        self.assertEqual(events[0]['request_id'], events[1]['request_id'])
        self.assertEqual(403, events[1]['http_status'])
        self.assertEqual('detail1', events[1]['instance_id'])
        self.assertNotIn('real-test-key', json.dumps(events))

    def test_logger_retains_usage_and_long_image_prompt_and_redacts_echoed_key(self):
        events = []
        logger = APIDebugLogger(on_event=events.append)
        logger.secrets = ['sample-secret']
        logger.log('response', response_json={'usage': {'total_tokens': 123}, 'image_prompt_en': 'a' * 4000, 'message': 'sample-secret'})
        self.assertEqual(123, events[0]['response_json']['usage']['total_tokens'])
        self.assertEqual(4000, len(events[0]['response_json']['image_prompt_en']))
        self.assertNotIn('sample-secret', json.dumps(events))

    def test_typesetting_follows_right_copy_box_and_preserves_product_area(self):
        with tempfile.TemporaryDirectory() as temp:
            source, output = Path(temp)/'source.png', Path(temp)/'output.jpg'
            Image.new('RGB', (1024, 1024), 'white').save(source)
            create_german_composite(source, output, 'Reise organisiert\nOrdnung im Koffer\nDetails im Blick\nPlatz sinnvoll nutzen', creative_plan={'pc': deepcopy(LAYOUTS['right'])})
            with Image.open(output) as image:
                self.assertGreater(image.getpixel((100, 500))[0], 245)
                self.assertLess(image.getpixel((620, 500))[0], 150)


class DesktopTests(unittest.TestCase):
    def test_upload_previews_entry_text_undo_plan_edit_and_api_inspector(self):
        with tempfile.TemporaryDirectory() as temp:
            root = tk.Tk()
            root.withdraw()
            try:
                app = AmazonImageBriefApp(root, Path(temp))
                root.update_idletasks()
                first, second = Path(temp)/'first.png', Path(temp)/'second.png'
                Image.new('RGB', (300, 300), 'blue').save(first)
                Image.new('RGB', (600, 500), 'red').save(second)
                with patch('amazon_image_brief.gui.filedialog.askopenfilenames', side_effect=[(str(first),), (str(first), str(second))]):
                    app.pick_product_images()
                    app.pick_product_images()
                self.assertEqual([str(first), str(second)], app.product_images)
                self.assertEqual(2, len(app.product_gallery.tree.get_children()))
                self.assertIsNotNone(app.product_gallery.preview_photo)
                entry = next(widget for widget in app.input_undo.histories if not isinstance(widget, tk.Text) and str(widget.cget('textvariable')) == str(app.variant_vars['color']))
                app.variant_tree.selection_set('0')
                app._load_variant_editor()
                original = entry.get()
                event = SimpleNamespace(widget=entry, type=tk.EventType.KeyPress, keysym='a', state=0)
                app.input_undo.before_edit(event)
                app.variant_vars['color'].set('Black')
                app.input_undo.undo(event)
                self.assertEqual(original, entry.get())
                app._autosave_variant()
                self.assertEqual(original, app.variants[0].color)
                app.input_undo.redo(event)
                self.assertEqual('Black', entry.get())
                app.variant_tree.selection_set('1')
                app._load_variant_editor()
                app.input_undo.undo(event)
                self.assertEqual('', entry.get())
                editor = app.current_copy_text
                app._set_text(editor, 'original', editable=True)
                editor.edit_separator()
                editor.insert('end', ' changed')
                app.input_undo.restore(editor, False)
                self.assertEqual('original', editor.get('1.0', 'end-1c'))
                prompt_ids = app.confirm_tree.get_children()
                app.confirm_tree.selection_set(prompt_ids[0])
                app._on_confirm_select()
                app.module_prompt_text.insert('end', 'old module edit')
                app.confirm_tree.selection_set(prompt_ids[1])
                app._on_confirm_select()
                app.input_undo.restore(app.module_prompt_text, False)
                self.assertEqual('', app.module_prompt_text.get('1.0', 'end-1c'))
                app.preview_creative_plan()
                self.assertEqual(len(app.module_instances), len(app.creative_workspace.tree.get_children()))
                app.api_inspector.add_event({'phase': 'request', 'request_id': 'one', 'method': 'POST', 'url': 'https://example.test', 'request_json': {'test': True}})
                app.api_inspector.add_event({'phase': 'response', 'request_id': 'one', 'http_status': 200, 'response_json': {'ok': True}})
                self.assertIn('200', app.api_inspector.response_text.get('1.0', 'end'))
                self.assertEqual(4, len(app.copy_workspace_tabs.tabs()))
            finally:
                root.destroy()
