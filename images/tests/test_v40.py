from copy import deepcopy
from pathlib import Path
import json
import tempfile
import threading
import time
import tkinter as tk
import unittest
from unittest.mock import patch
import zipfile

from PIL import Image
from openpyxl import load_workbook

from amazon_image_brief.direct_workspace import DirectImageApp
from amazon_image_brief.direct_images import module_recipe, append_assets, image_jobs, validate_recipe, DirectImageService, export_direct
from amazon_image_brief.models import ProductProject, ModuleInstance, GenerationOptions
from amazon_image_brief.ai_client import OpenAIClient, OpenAIError
from amazon_image_brief.generation_control import GenerationControl, GenerationCancelled
from amazon_image_brief.profile_store import ProfileStore


def sample(root):
    product = root/'product.png'
    reference = root/'reference.png'
    Image.new('RGB', (800, 800), 'blue').save(product)
    Image.new('RGB', (800, 600), 'yellow').save(reference)
    project = ProductProject(project_name='Direct tests', brand='A', product_name_zh='行李箱', selling_points='双排轮',
                             module_instances=[ModuleInstance('m1', 'MAIN_DETAIL', '主图'), ModuleInstance('m2', 'MAIN_SCENE', '主图')],
                             copy_languages=['de'], options=GenerationOptions(output_root=str(root/'outputs'), api_debug_enabled=False))
    for instance in project.module_instances:
        recipe = module_recipe(project, instance.instance_id)
        recipe = append_assets(recipe, 'product', [str(product)])
        recipe = append_assets(recipe, 'reference', [str(reference)])
        recipe['direct_prompt'] = '以 P01 为产品，参考 R01 的色调和文案位置；保持本体。'
        project.module_recipes[instance.instance_id] = recipe
    return project


class ImageOnlyClient:
    image_available = True
    available = False  # Text configuration is intentionally missing.
    def __init__(self, fail_language='', hook=None):
        self.debug_context = {}
        self.calls = []
        self.fail_language, self.hook = fail_language, hook
        self.options = GenerationOptions(image_provider='openai', image_model='gpt-image-2')
    def generate_json(self, *_a, **_k):
        raise AssertionError('No copy/context/layout endpoint may be called')
    analyze_images_json = generate_json
    def generate_image(self, prompt, destination, **kwargs):
        self.calls.append((prompt, deepcopy(kwargs), deepcopy(self.debug_context)))
        if self.hook:
            self.hook()
        if f'({self.fail_language})' in prompt and self.fail_language:
            raise OpenAIError('Mock language failure')
        destination.parent.mkdir(parents=True, exist_ok=True)
        Image.new('RGB', (900, 900), 'green').save(destination, 'JPEG')


class DirectServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.project = sample(self.root)
    def tearDown(self): self.temp.cleanup()

    def test_default_german_and_exact_module_times_language_count(self):
        self.assertEqual(['de', 'de'], [lang for _, lang, _ in image_jobs(self.project)])
        self.project.copy_languages = ['de', 'fr', 'de']
        jobs = image_jobs(self.project)
        self.assertEqual(4, len(jobs))
        self.assertNotIn('zh', [lang for _, lang, _ in jobs])
        self.project.copy_languages = []
        with self.assertRaisesRegex(ValueError, '至少勾选'):
            image_jobs(self.project)

    def test_white_main_still_creates_each_language_but_no_text(self):
        self.project.module_instances[0].module_code = 'MAIN_WHITE'
        self.project.copy_languages = ['de', 'fr']
        client = ImageOnlyClient()
        result = DirectImageService(self.root, client).run(self.project, ids=['m1'])
        self.assertEqual(2, result['completed'])
        self.assertTrue(all('MANDATORY MAIN_WHITE' in prompt for prompt, _, _ in client.calls))

    def test_optional_reference_and_empty_prompt_validation(self):
        recipe = self.project.module_recipes['m1']
        recipe['direct_reference_images'] = []
        recipe['direct_prompt'] = '以 P01 为实物，白色机场场景。'
        validate_recipe(recipe, '模块')
        recipe['direct_prompt'] = '  '
        with self.assertRaisesRegex(ValueError, 'Prompt'):
            validate_recipe(recipe, '模块')

    def test_real_image_adapter_wire_request_contains_numbered_inputs_only(self):
        from test_v302 import ImageResponse
        events = []
        options = GenerationOptions(image_provider='openai', image_protocol='openai_images', image_model='gpt-image-2',
                                    cache_enabled=False, max_retries=0, requests_per_minute=100000, api_debug_enabled=False)
        client = OpenAIClient(options=options, image_api_key='unit-v40-image-key', on_api_event=events.append)
        with patch('urllib.request.urlopen', return_value=ImageResponse()) as network:
            result = DirectImageService(self.root, client).run(self.project, ids=['m1'])
        self.assertEqual(1, result['completed'])
        self.assertEqual(1, network.call_count)
        request = network.call_args.args[0]
        self.assertTrue(request.full_url.endswith('/images/edits'))
        self.assertEqual(2, request.data.count(b'name="image[]"'))
        self.assertIn(b'P01', request.data)
        self.assertIn(b'R01', request.data)
        self.assertNotIn('unit-v40-image-key', json.dumps(events))
        self.assertTrue(any(e.get('output_language') == 'de' for e in events))

    def test_append_delete_keeps_stable_ids_and_independent_roles(self):
        recipe = self.project.module_recipes['m1']
        p2 = self.root/'p2.jpg'
        Image.new('RGB', (500, 500)).save(p2)
        updated = append_assets(recipe, 'product', [recipe['direct_product_images'][0]['asset_path'], str(p2)])
        self.assertEqual(['P01', 'P02'], [r['id'] for r in updated['direct_product_images']])
        updated['direct_product_images'] = updated['direct_product_images'][1:]
        updated = append_assets(updated, 'product', [str(self.root/'product.png')])
        self.assertEqual(['P02', 'P03'], [r['id'] for r in updated['direct_product_images']])
        self.assertEqual(['R01'], [r['id'] for r in updated['direct_reference_images']])
        with self.assertRaisesRegex(ValueError, 'P01'):
            validate_recipe(updated, '测试模块')
        self.assertEqual(1, len(self.project.module_recipes['m2']['direct_product_images']))

    def test_direct_image_request_has_exact_prompt_roles_ids_language_no_copy_api(self):
        self.project.copy_languages = ['de', 'fr']
        client, events = ImageOnlyClient(), []
        result = DirectImageService(self.root, client, events.append).run(self.project)
        self.assertEqual(4, result['completed'])
        self.assertEqual(4, len(client.calls))
        prompt, kwargs, debug = client.calls[0]
        self.assertIn(self.project.module_recipes['m1']['direct_prompt'], prompt)
        self.assertIn('OUTPUT LANGUAGE: 德语 (de)', prompt)
        self.assertIn('PRODUCT IDENTITY LOCK', prompt)
        self.assertIn('reference_id', prompt)
        self.assertNotIn('APPROVED VISUAL PLAN', prompt)
        self.assertEqual(2, len(kwargs['reference_images']))
        self.assertEqual(['P01', 'R01'], [r['reference_id'] for r in debug['reference_inputs']['sources']])
        ready = [e for e in events if e['result']['status'] == '已完成']
        self.assertEqual(4, len(ready))
        self.assertEqual(1, ready[0]['done'])
        self.assertEqual(4, ready[0]['total'])
        with Image.open(ready[0]['result']['final_image']) as image:
            self.assertEqual('JPEG', image.format)

    def test_single_image_gateway_contact_sheet_keeps_number_mapping(self):
        client = ImageOnlyClient()
        client.options.image_provider = 'custom'
        client.options.image_protocol = 'custom_http'
        DirectImageService(self.root, client).run(self.project, ids=['m1'])
        _, kwargs, debug = client.calls[0]
        self.assertEqual(1, len(kwargs['reference_images']))
        self.assertIn('contact-sheet', debug['reference_inputs']['strategy'])
        self.assertEqual(['P01', 'R01'], [r['reference_id'] for r in debug['reference_inputs']['sources']])

    def test_reference_cannot_replace_product_and_bad_input_sends_no_request(self):
        self.project.module_recipes['m2']['direct_product_images'] = []
        client = ImageOnlyClient()
        with self.assertRaisesRegex(ValueError, '本体图'):
            DirectImageService(self.root, client).run(self.project)
        self.assertFalse(client.calls)

    def test_language_failure_does_not_block_other_languages(self):
        self.project.copy_languages = ['de', 'fr', 'it']
        events = []
        result = DirectImageService(self.root, ImageOnlyClient('fr'), events.append).run(self.project, ids=['m1'])
        self.assertEqual((2, 1), (result['completed'], result['failed']))
        self.assertEqual('生成失败', [e for e in events if e['language'] == 'fr'][-1]['result']['status'])

    def test_regenerate_one_language_archives_original_and_skips_other_language(self):
        events = []
        DirectImageService(self.root, ImageOnlyClient(), events.append).run(self.project, ids=['m1'])
        old = events[-1]['result']
        self.project.module_recipes['m1']['direct_results']['de'] = old
        original_bytes = Path(old['final_image']).read_bytes()
        client = ImageOnlyClient()
        DirectImageService(self.root, client, events.append).run(self.project, ids=['m1'], languages=['de'], revisions={('m1', 'de'): '背景改为浅灰'})
        current = events[-1]['result']
        self.assertEqual(2, current['version'])
        self.assertEqual(old['final_image'], current['history'][0]['final_image'])
        self.assertEqual(original_bytes, Path(old['final_image']).read_bytes())
        self.assertIn('背景改为浅灰', client.calls[0][0])
        self.assertFalse(client.calls[0][1]['use_cache'])
        self.assertEqual(1, len(client.calls))

    def test_pause_displays_returned_image_then_resumes(self):
        control = GenerationControl('image')
        control.start()
        events, errors = [], []
        client = ImageOnlyClient(hook=lambda: control.pause() if not events else None)
        # Pause after the first request has been sent, not before its result.
        client.hook = lambda: control.pause() if len(client.calls) == 1 else None
        def run():
            try: DirectImageService(self.root, client, events.append, control.checkpoint).run(self.project)
            except Exception as exc: errors.append(exc)
        thread = threading.Thread(target=run)
        thread.start()
        deadline = time.monotonic()+10
        while time.monotonic() < deadline and control.state != 'paused': time.sleep(.01)
        try:
            self.assertEqual('paused', control.state)
            self.assertEqual('已完成', events[-1]['result']['status'])
            self.assertEqual(1, len(client.calls))
        finally:
            control.resume()
            thread.join(10)
            control.close()
        self.assertEqual([], errors)
        self.assertFalse(thread.is_alive())
        self.assertEqual(2, len(client.calls))

    def test_stop_drops_late_return_and_no_following_calls(self):
        control = GenerationControl('image')
        control.start()
        client, events = ImageOnlyClient(hook=control.stop), []
        with self.assertRaises(GenerationCancelled):
            DirectImageService(self.root, client, events.append, control.checkpoint).run(self.project)
        self.assertEqual(1, len(client.calls))
        self.assertFalse(any(e['result']['status'] == '已完成' for e in events))

    def test_export_only_chosen_languages_and_profile_assets(self):
        self.project.copy_languages = ['de', 'fr']
        def save(event):
            self.project.module_recipes[event['id']]['direct_results'][event['language']] = event['result']
        DirectImageService(self.root, ImageOnlyClient(), save).run(self.project)
        result = export_direct(self.project, self.root/'export')
        workbook = load_workbook(result['excel_path'])
        self.assertEqual(5, workbook.active.max_row)
        self.assertEqual(4, len(workbook.active._images))
        with zipfile.ZipFile(result['archive_path']) as archive:
            self.assertEqual(4, len(archive.namelist()))
            self.assertTrue(all(name.endswith('.jpg') for name in archive.namelist()))
        store = ProfileStore(self.root/'data')
        identity = store.save_as({'project': self.project.to_dict(), 'ui': {}}, 'Test')
        restored = store.load(identity)['project']['module_recipes']['m1']
        self.assertEqual('P01', restored['direct_product_images'][0]['id'])
        self.assertIn('assets', restored['direct_product_images'][0]['asset_path'])
        self.assertTrue(Path(restored['direct_results']['de']['final_image']).is_file())


class DirectDesktopTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.models = patch.object(OpenAIClient, 'list_models', return_value=['unit-image'])
        self.models.start()
        self.root = tk.Tk()
        self.root.withdraw()
        self.app = DirectImageApp(self.root, self.folder)
        self.app.populate_project(sample(self.folder))
        self.root.update_idletasks()
    def tearDown(self):
        for control in self.app.generation_controls.values(): control.close()
        pool = getattr(self.root, '_image_preview_pool', None)
        if pool:
            pool.close()
            for worker in pool._threads:
                worker.join(5)
        for callback in self.root.tk.call('after', 'info'): self.root.after_cancel(callback)
        self.root.destroy()
        self.models.stop()
        self.temp.cleanup()

    def test_only_one_generation_page_and_no_layout_or_copy_tabs_visible(self):
        app = self.app
        labels = [app.tabs.tab(tab, 'text') for tab in app.tabs.tabs()]
        self.assertEqual(5, len(labels))
        self.assertIn('04  创意生图', labels)
        self.assertNotIn('05  文案生成', labels)
        self.assertFalse(app.copy_workspace_tabs.winfo_ismapped())
        self.assertEqual(3, len(app.direct_tabs.tabs()))
        project = app.collect_project(sync_modules=False)
        self.assertFalse(project.options.optimize_copy_with_ai)
        self.assertFalse(project.options.ai_plan_before_copy)
        self.assertFalse(project.options.context_before_generation)
        self.assertEqual(['de'], project.copy_languages)

    def test_module_inputs_are_separate_and_append_preview_batch_delete(self):
        app = self.app
        self.assertEqual('m1', app.direct_current_id)
        added = self.folder/'new.jpg'
        Image.new('RGB', (600, 600)).save(added)
        app._append_direct_assets('product', [str(added), str(added)])
        gallery = app.direct_galleries['product']
        self.assertEqual(2, len(gallery.paths))
        self.assertEqual('P02', gallery.tree.set('1', 'number'))
        deadline = time.monotonic()+5
        while gallery.preview_photo is None and time.monotonic() < deadline:
            self.root.update()
            time.sleep(.01)
        self.assertIsNotNone(gallery.preview_photo)
        gallery.toggle_checked('0')
        gallery.remove_selected()
        self.assertEqual(['P02'], [r['id'] for r in app.module_recipes['m1']['direct_product_images']])
        self.assertEqual(1, len(app.module_recipes['m1']['direct_reference_images']))
        self.assertEqual(['P01'], [r['id'] for r in app.module_recipes['m2']['direct_product_images']])

    def test_prompt_autosave_switch_and_profile_restore(self):
        app = self.app
        app._set_text(app.direct_prompt_text, '以 P01 为准，R01 杂志风格', editable=True)
        app.direct_module_tree.selection_set('m2')
        app._select_direct_module()
        self.assertEqual('以 P01 为准，R01 杂志风格', app.module_recipes['m1']['direct_prompt'])
        app.language_vars['fr'].set(True)
        app._direct_languages_changed()
        snapshot = app.profiles.capture()
        app.profiles.restore(snapshot)
        self.assertEqual('m2', app.direct_current_id)
        self.assertEqual('以 P01 为准，R01 杂志风格', app.module_recipes['m1']['direct_prompt'])
        self.assertIn('4 张图', app.direct_count_var.get())

    def test_each_return_updates_language_preview_before_all_done(self):
        app = self.app
        app.busy = True
        path = self.folder/'result.jpg'
        Image.new('RGB', (800, 800)).save(path)
        event = dict(id='m1', language='de', done=1, total=4, output_dir=str(self.folder),
                     result=dict(status='已完成', final_image=str(path), version=1))
        app._receive_direct_image(event)
        self.assertEqual(str(path), app.direct_preview_paths['after'])
        deadline = time.monotonic()+5
        while 'after' not in app.direct_photos and time.monotonic() < deadline:
            self.root.update()
            time.sleep(.01)
        self.assertIsNotNone(app.direct_photos['after'])
        self.assertEqual(str(app.direct_results_tab), app.direct_tabs.select())
        self.assertIn('1/4', app.progress_var.get())
        app.busy = False

    def test_generation_does_not_require_text_key_and_uses_module_snapshot(self):
        app = self.app
        app.language_vars['fr'].set(True)
        app._direct_languages_changed()
        client = ImageOnlyClient()
        captured = []
        with patch.object(app, '_build_ai_client', return_value=client), patch('tkinter.messagebox.askyesno', return_value=True), patch.object(app, '_start_async', side_effect=lambda op, fn, msg: captured.append((op, fn, msg))):
            app.generate_direct('module')
        self.assertEqual('direct_images', captured[0][0])
        self.assertIn('2 张', captured[0][2])
        result = captured[0][1]()
        self.assertEqual(2, result['completed'])
        self.assertEqual(2, len(client.calls))

    def test_duplicate_module_copies_inputs_but_not_old_generated_results(self):
        app = self.app
        app.module_recipes['m1']['direct_results'] = {'de': {'status': '已完成', 'final_image': 'old.jpg'}}
        app.confirm_tree.selection_set('m1')
        before = set(app.module_order)
        app.duplicate_module_instance(app.confirm_tree)
        added = (set(app.module_order)-before).pop()
        self.assertEqual({}, app.module_recipes[added]['direct_results'])
        self.assertEqual('P01', app.module_recipes[added]['direct_product_images'][0]['id'])

    def test_small_window_can_scroll_both_input_galleries_and_results(self):
        app = self.app
        self.root.attributes('-alpha', 0)
        self.root.geometry('1120x720')
        self.root.deiconify()
        app.tabs.select(app.direct_page)
        app.direct_tabs.select(app.direct_inputs_tab)
        self.root.update()
        app.direct_inputs_tab.scroll_canvas.yview_moveto(1)
        self.root.update()
        self.assertAlmostEqual(1, app.direct_inputs_tab.scroll_canvas.yview()[1], places=2)
        self.assertGreater(app.direct_galleries['reference'].winfo_width(), 500)
        app.direct_tabs.select(app.direct_results_tab)
        self.root.update()
        app.direct_results_tab.scroll_canvas.yview_moveto(1)
        self.assertAlmostEqual(1, app.direct_results_tab.scroll_canvas.yview()[1], places=2)

    def test_async_buttons_route_results_to_unified_ui_and_finish(self):
        app, client = self.app, ImageOnlyClient()
        app.language_vars['fr'].set(True)
        app._direct_languages_changed()
        with patch.object(app, '_build_ai_client', return_value=client), patch('tkinter.messagebox.askyesno', return_value=True), patch('tkinter.messagebox.showerror') as errors:
            app.generate_direct('all')
            deadline = time.monotonic()+15
            while app.busy and time.monotonic() < deadline:
                self.root.update()
                time.sleep(.01)
            self.assertFalse(app.busy)
            errors.assert_not_called()
        self.assertEqual(4, len(client.calls))
        self.assertEqual('m2', app.direct_current_id)
        self.assertEqual('fr', app.direct_current_language)
        self.assertIn('成功 4', app.status_var.get())
        for iid in ('m1', 'm2'):
            self.assertEqual({'de', 'fr'}, set(app.module_recipes[iid]['direct_results']))
        self.assertEqual('normal', str(app.direct_prompt_text.cget('state')))

    def test_returned_preview_is_scrolled_into_small_window(self):
        app = self.app
        self.root.attributes('-alpha', 0)
        self.root.geometry('1120x720')
        self.root.deiconify()
        self.root.update()
        path = self.folder/'preview.jpg'
        Image.new('RGB', (800, 800)).save(path)
        app.busy = True
        app._receive_direct_image(dict(id='m1', language='de', done=1, total=2, output_dir=str(self.folder),
                                        result=dict(status='已完成', final_image=str(path), version=1)))
        self.root.update()
        canvas, preview = app.direct_results_tab.scroll_canvas, app.direct_preview_labels['after']
        self.assertGreaterEqual(preview.winfo_rooty(), canvas.winfo_rooty())
        self.assertLess(preview.winfo_rooty()+preview.winfo_height(), canvas.winfo_rooty()+canvas.winfo_height()+15)
        app.busy = False


if __name__ == '__main__': unittest.main()
