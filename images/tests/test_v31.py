from copy import deepcopy
from pathlib import Path
from io import BytesIO
import json
import os
import tempfile
import threading
import time
import tkinter as tk
from tkinter import ttk
import unittest
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook

from amazon_image_brief.gui import AmazonImageBriefApp
from amazon_image_brief.models import ProductProject, GenerationOptions
from amazon_image_brief.brief_generator import BriefGenerator
from amazon_image_brief.product_context import ProductContext, context_fingerprint, product_facts
from amazon_image_brief.service import GenerationService
from amazon_image_brief.ai_client import OpenAIClient
from amazon_image_brief.env_store import read_env_values, provider_key_name
from amazon_image_brief.generation_control import GenerationControl, GenerationCancelled
from amazon_image_brief.languages import TARGET_LANGUAGES
from amazon_image_brief.costing import estimate_project_cost
from test_v28 import PlannedAI, project_at
from test_v30 import ContextAI, photos


class LocaleBrandTests(unittest.TestCase):
    def test_new_default_and_legacy_roundtrip(self):
        self.assertEqual(('de', 'zh'), ProductProject().active_languages())
        self.assertTrue({'pt', 'nl', 'pl', 'fi', 'el', 'ro', 'uk', 'en'} <= set(TARGET_LANGUAGES))
        raw = {'copy_overrides': {'one': {'en': 'English', 'de': 'German', 'fr': 'French', 'zh': '中文'}}}
        legacy = ProductProject.from_dict(raw)
        self.assertEqual({'en', 'de', 'fr', 'zh'}, set(legacy.active_languages()))
        legacy.brand_slogan = 'Travel freely'
        self.assertEqual(legacy.to_dict(), ProductProject.from_dict(legacy.to_dict()).to_dict())

    def test_selected_languages_and_per_language_chinese_generated(self):
        project = project_at('unused')
        project.copy_languages = ['de', 'nl', 'pl']
        project.options.optimize_copy_with_ai = True
        client = PlannedAI()
        briefs = BriefGenerator(client).generate(project)
        self.assertEqual({'de', 'nl', 'pl', 'zh'}, set(briefs[1].copy))
        self.assertEqual({'de', 'nl', 'pl'}, set(briefs[1].chinese_translations))
        self.assertIn('AI完成', briefs[1].generation_status)
        self.assertIn('chinese_translations', client.requests[0])
        self.assertIn('nl', client.requests[0])

    def test_missing_chinese_backtranslation_fails_validation(self):
        project = project_at('unused')
        brief = BriefGenerator()._local_briefs(project)[1]
        prompt = BriefGenerator()._copy_prompt(project, brief)
        result = PlannedAI().generate_json(prompt, 'test')
        result['chinese_translations'].pop('de')
        with self.assertRaisesRegex(Exception, '逐行中文对照'):
            BriefGenerator._validate_copy_result(brief, result, [])

    def test_logo_is_analyzed_as_brand_not_product_and_invalidates_context(self):
        with tempfile.TemporaryDirectory() as temp:
            project = project_at(temp)
            project.options.context_before_generation = True
            paths = photos(Path(temp), 2)
            project.product_image_paths = paths[:1]
            project.logo_image_path = paths[1]
            project.brand_slogan = 'Travel freely'
            project.brand_colors = '#123456'
            client = ContextAI()
            ProductContext(client).prepare(project)
            self.assertEqual(2, len(client.calls))
            self.assertIn('BRAND LOGO ONLY', client.calls[1][1])
            self.assertIn('Travel freely', client.calls[0][1])
            self.assertEqual('#123456', product_facts(project)['brand_colors'])
            before = context_fingerprint(project)
            Image.new('RGB', (800, 1000), 'yellow').save(paths[1])
            self.assertNotEqual(before, context_fingerprint(project))

    def test_image_uses_selected_exact_language(self):
        project = project_at('unused')
        project.copy_languages = ['de', 'nl']
        project.options.image_language = 'nl'
        project.brand_slogan = 'Travel freely'
        brief = BriefGenerator()._local_briefs(project)[1]
        brief.copy.update(nl='Exact Nederlands\nTweede regel', de='NICHT DRUCKEN')
        prompt = GenerationService._generation_prompt(project, brief)
        self.assertIn('Exact Nederlands\nTweede regel', prompt)
        self.assertNotIn('NICHT DRUCKEN', prompt)
        self.assertIn('Travel freely', prompt)

    def test_image_and_excel_are_independent_with_dynamic_columns(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            project = project_at(root/'out')
            project.copy_languages = ['de', 'nl', 'pl']
            service = GenerationService(root, ContextAI(), export_excel_on_images=False)
            result = service.run(project)
            self.assertEqual('', result['excel'])
            self.assertFalse(list((root/'out').rglob('*.xlsx')))
            with patch.object(service.ai_client, 'generate_json', side_effect=AssertionError('Excel must not call AI')):
                exported = service.refresh_output(project, result['briefs'], Path(result['output_dir']), export_excel=True)
            workbook = load_workbook(exported['excel'])
            headers = [cell.value for cell in workbook['主图'][3]]
            self.assertIn('荷兰语文案', headers)
            self.assertIn('波兰语文案', headers)
            self.assertIn('中文文案', headers)
            self.assertNotIn('法语', headers)
            self.assertIn('语言与中文对照', workbook.sheetnames)
            workbook.close()

    def test_cost_increases_with_language_selection(self):
        project = project_at('unused')
        project.options.optimize_copy_with_ai = True
        project.copy_languages = ['de']
        small = estimate_project_cost(project).output_tokens
        project.copy_languages = list(TARGET_LANGUAGES)
        self.assertGreater(estimate_project_cost(project).output_tokens, small)

    def test_cancelled_response_is_not_saved_as_image(self):
        from test_v302 import ImageResponse
        with tempfile.TemporaryDirectory() as temp:
            control = GenerationControl('image')
            control.start()
            client = OpenAIClient(options=GenerationOptions(api_debug_enabled=False), image_api_key='test-only')
            client.request_control = control.checkpoint
            def response(*_args, **_kwargs):
                control.stop()
                return ImageResponse()
            with patch('urllib.request.urlopen', response):
                with self.assertRaises(GenerationCancelled):
                    client._post_json('https://example.test', '/images', {}, 'test-only', '', 'test')
            self.assertFalse(list(Path(temp).iterdir()))


class DesktopV31Tests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {}, clear=False)
        self.env.start()
        self.models = patch.object(OpenAIClient, 'list_models', return_value=['live-model-1', 'live-model-2'])
        self.models.start()
        self.root = tk.Tk()
        self.root.withdraw()
        self.app = AmazonImageBriefApp(self.root, Path(self.temp.name))
        self.root.update_idletasks()

    def tearDown(self):
        for control in self.app.generation_controls.values():
            control.close()
        for callback in self.root.tk.call('after', 'info'):
            self.root.after_cancel(callback)
        self.root.destroy()
        self.models.stop()
        self.env.stop()
        self.temp.cleanup()

    def pump(self, condition=lambda: True, timeout=5):
        deadline = time.monotonic()+timeout
        while time.monotonic() < deadline:
            self.root.update()
            if condition():
                return
            time.sleep(.01)
        self.fail('UI callback timeout')

    def select_provider(self, kind, provider):
        self.app.ai_config_vars[kind]['provider'].set(self.app._provider_presets(kind)[provider].label)
        self.app.apply_provider_preset(kind)

    def test_named_page_order_merged_prompt_and_scrollbars(self):
        app = self.app
        self.assertEqual(['API配置', '产品信息', '竞品与关键词', '选择模块', '文案生成', '图片生成', '批量任务与合规'],
                         [app.tabs.tab(tab, 'text').split('  ', 1)[1] for tab in app.tabs.tabs()])
        self.assertEqual(app.pages['api'], app.tabs.select())
        self.assertIs(app.confirm_tree, app.module_selected_tree)
        self.assertEqual(3, len(app.export_workspace_tabs.tabs()))
        self.assertEqual(app.image_review_body, app.recipe_panels['image']['frame'].master)
        self.assertTrue(app.current_copy_text.cget('yscrollcommand'))
        self.assertTrue(app.copy_scroll_frame.scroll_canvas.cget('yscrollcommand'))
        self.assertEqual('clam', ttk.Style(self.root).theme_use())
        self.assertTrue(app.language_vars['de'].get())
        self.assertEqual(1, sum(var.get() for var in app.language_vars.values()))

    def test_named_configuration_slogan_languages_and_logo_persist(self):
        app = self.app
        image = Path(self.temp.name)/'logo.png'
        Image.new('RGB', (100, 60), '#123456').save(image)
        app.logo_image_path = str(image)
        app.scalar_vars['brand_colors'].set('#123456 / #ffffff')
        app.scalar_vars['brand_slogan'].set('Travel freely')
        app.language_vars['nl'].set(True)
        app._languages_changed()
        app._refresh_brand_preview()
        self.assertIsNotNone(app._brand_photo)
        self.assertEqual(2, len(app.brand_swatches.winfo_children()))
        app.profiles.name_var.set('旅行箱 DE + NL')
        app.profiles.save_current()
        identity = app.profiles.profile_id
        self.assertTrue(identity)
        app.profiles.name_var.set('新的配置名称')
        app.profiles.save_current()
        saved = app.profiles.store.load(identity)
        self.assertEqual('新的配置名称', saved['name'])
        self.assertEqual('Travel freely', saved['project']['brand_slogan'])
        self.assertEqual({'de', 'nl'}, set(saved['project']['copy_languages']))
        app.profiles.restore(saved)
        self.assertTrue(Path(app.logo_image_path).is_file())
        self.assertEqual('Travel freely', app.scalar_vars['brand_slogan'].get())

    def test_provider_keys_are_separate_and_fallback_fetches_models(self):
        app = self.app
        for kind in ('text', 'image'):
            self.select_provider(kind, 'cunai')
            app.ai_key_vars[kind].set('test-cun-'+kind)
            self.select_provider(kind, 'sudocode')
            app.ai_key_vars[kind].set('test-sudo-'+kind)
            self.select_provider(kind, 'cunai')
            self.assertEqual('test-cun-'+kind, app.ai_key_vars[kind].get())
            getattr(app, kind+'_fallback_provider_var').set('sudocode')
            app._apply_fallback_provider(kind)
        self.pump(lambda: app.image_fallback_model_var.get() == 'live-model-1' and app.text_fallback_model_var.get() == 'live-model-1')
        values = read_env_values(Path(self.temp.name)/'.env')
        self.assertEqual('test-cun-text', values[provider_key_name('text', 'cunai')])
        app.auto_failover_var.set(True)
        client = app._build_ai_client(app.collect_project().options)
        self.assertEqual('test-sudo-text', client.text_fallbacks[0].text_api_key)
        self.assertEqual('test-sudo-image', client.image_fallbacks[0].image_api_key)
        document = json.dumps(app.profiles.capture())
        self.assertNotIn('test-cun-', document)
        self.assertNotIn('test-sudo-', document)

    def test_late_model_list_cannot_replace_new_provider(self):
        app = self.app
        app.text_fallback_provider_var.set('sudocode')
        app.text_fallback_model_var.set('keep-model')
        app._model_request_ids[('text', True)] = 2
        app._model_events.put(('models', (('text', True), 1, 'cunai', ['wrong'], '')))
        app._poll_model_events()
        self.assertEqual('keep-model', app.text_fallback_model_var.get())

    def test_stopping_task_fences_late_results_and_new_task_works(self):
        app = self.app
        started, release, returned = threading.Event(), threading.Event(), threading.Event()
        def old_work():
            started.set()
            release.wait(4)
            returned.set()
            return 'OLD'
        handled = []
        with patch.object(app, '_handle_operation_done', side_effect=lambda op, value: handled.append(value)), patch('amazon_image_brief.generation_workspace.messagebox.askyesno', return_value=True):
            app._start_async('draft', old_work, 'test')
            self.assertTrue(started.wait(1))
            previous = app.generation_controls['copy']
            app.stop_generation()
            self.assertFalse(app.busy)
            with self.assertRaises(GenerationCancelled):
                previous.checkpoint()
            self.assertEqual('disabled', str(app.control_widgets['copy'][1]['state']))
            app._start_async('draft', lambda: 'NEW', 'new test')
            release.set()
            self.pump(lambda: returned.is_set() and not app.busy)
            self.assertEqual(['NEW'], handled)

    def test_edit_invalidates_old_chinese_comparison(self):
        app = self.app
        project = project_at(self.temp.name)
        project.options.optimize_copy_with_ai = True
        app.populate_project(project)
        app.current_briefs = BriefGenerator(PlannedAI()).generate(project)
        app._refresh_review_tree()
        app.review_tree.selection_set('detail1')
        app._on_review_select()
        app._set_text(app.current_copy_text, 'Changed German', editable=True)
        app._commit_review_edits()
        app._refresh_language_compare()
        brief = app._brief_by_code('detail1')
        self.assertNotIn('de', brief.chinese_translations)
        self.assertIn('暂无逐语回译', app.translation_compare_text.get('1.0', 'end'))

    def test_small_window_can_scroll_copy_and_image_review_to_bottom(self):
        app = self.app
        self.root.attributes('-alpha', 0)
        self.root.geometry('1120x720')
        self.root.deiconify()
        app.tabs.select(app.pages['copy'])
        self.root.update()
        canvas = app.copy_scroll_frame.scroll_canvas
        self.assertLess(canvas.yview()[1], 1)
        canvas.yview_moveto(1)
        self.assertGreater(canvas.yview()[0], 0)
        app.tabs.select(app.pages['image'])
        self.root.update()
        outer = app.export_scroll_frame.scroll_canvas
        self.assertLess(outer.yview()[1], 1)
        outer.yview_moveto(1)
        self.assertGreater(outer.yview()[0], 0)
        inner = app.image_review_page.scroll_canvas
        self.assertLess(inner.yview()[1], 1)
        inner.yview_moveto(1)
        self.assertGreater(inner.yview()[0], 0)
        self.assertTrue(app.control_widgets['image'][3].winfo_exists())


if __name__ == '__main__':
    unittest.main()
