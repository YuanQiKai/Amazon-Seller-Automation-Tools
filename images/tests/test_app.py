from __future__ import annotations

import json
import tempfile
import time
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw

from amazon_image_brief.ai_client import OpenAIClient, OpenAIError
from amazon_image_brief.ai_runtime import AICache
from amazon_image_brief.batch_queue import PersistentBatchQueue
from amazon_image_brief.brief_generator import BriefGenerator
from amazon_image_brief.catalogs import default_module_selection
from amazon_image_brief.data_io import MarketWorkbookManager
from amazon_image_brief.costing import estimate_project_cost
from amazon_image_brief.excel_exporter import ExcelExporter
from amazon_image_brief.image_composer import create_german_composite, create_placeholder
from amazon_image_brief.models import GenerationOptions, Keyword, ModuleBrief, ModuleInstance, ProductProject, ProductVariant
from amazon_image_brief.palette import extract_logo_palette
from amazon_image_brief.presets import AMAZON_MARKETPLACES, IMAGE_MODELS, TEXT_MODELS
from amazon_image_brief.providers import (
    IMAGE_PROVIDER_PRESETS,
    PROVIDER_CONFIG_PATH,
    TEXT_PROVIDER_PRESETS,
    model_catalog_url,
)
from amazon_image_brief.quality import image_similarity, text_overflow_report
from amazon_image_brief.rules import RuleLibrary
from amazon_image_brief.service import GenerationService


class FakeAIClient:
    available = True

    def generate_json(self, prompt: str, model: str):
        if "Audit the following" in prompt:
            return [
                {
                    "module_code": "MAIN_DETAIL",
                    "verdict": "建议修改",
                    "issues": ["德语支持句略长"],
                    "suggestions": ["缩短支持句并保留核心卖点"],
                }
            ]
        if "Rewrite the copy for one" in prompt:
            return {
                "copy": {language: f"新版-{language}" for language in ("en", "zh", "de", "fr", "it", "es")},
                "design_brief_zh": "新版局部设计说明",
                "compliance_note_zh": "仅使用已验证信息",
                "image_prompt_en": "A refreshed no-text close-up composition.",
            }
        raise AssertionError("测试未覆盖的 JSON 请求")

    def generate_image(self, prompt, destination, model, quality, ratio, reference_images, **_kwargs):
        create_placeholder(Path(destination), "局部重生", prompt, ratio)


def sample_project(output_root: str, image_path: str = "") -> ProductProject:
    selected = default_module_selection()
    order = [code for channel in ("主图", "高级A+", "Brand Story", "品牌旗舰店") for code in selected[channel]]
    project = ProductProject(
        project_name="行李箱测试项目",
        brand="TestBrand",
        product_name_zh="20英寸旅行箱",
        product_name_en="20-inch Carry-On Suitcase",
        category="Luggage",
        marketplace="Amazon DE — 德国 (amazon.de)",
        material="Polycarbonate",
        functions="Expandable\nSpinner wheels",
        selling_points="Lightweight, organized storage, smooth rolling",
        variants=[
            ProductVariant("20寸", "CASE-20", "55×35×23cm", "Black", "2.8kg", "38L", "€89.99"),
            ProductVariant("24寸", "CASE-24", "65×43×28cm", "Silver", "3.7kg", "65L", "€109.99"),
        ],
        custom_fields={"轮子": "4组双排万向轮", "锁具": "TSA锁"},
        product_image_paths=[image_path] if image_path else [],
        keywords=[Keyword("carry on suitcase", "en", "category", 5)],
        selected_modules=selected,
        module_order=order,
        module_prompts={"MAIN_DETAIL": "Use a macro close-up of the double spinner wheel."},
        options=GenerationOptions(
            optimize_copy_with_ai=False,
            generate_ai_images=False,
            generate_german_composites=False,
            text_model="gpt-5.6-terra",
            image_model="gpt-image-2",
            output_root=output_root,
        ),
    )
    project.module_instances = project.normalized_module_instances()
    project.module_order = [item.instance_id for item in project.module_instances]
    return project


class CatalogAndModelTests(unittest.TestCase):
    def test_default_counts_and_marketplace_catalog(self) -> None:
        selected = default_module_selection()
        self.assertEqual(9, len(selected["主图"]))
        self.assertEqual(7, len(selected["高级A+"]))
        self.assertGreaterEqual(len(AMAZON_MARKETPLACES), 23)
        self.assertTrue(any("amazon.de" in item for item in AMAZON_MARKETPLACES))
        self.assertIn("gpt-5.6-terra", TEXT_MODELS)
        self.assertIn("gpt-image-2", IMAGE_MODELS)

    def test_project_round_trip_preserves_variants_and_review_state(self) -> None:
        project = sample_project("outputs")
        project.copy_versions = {"MAIN_DETAIL": [{"version": 1, "copy": {"de": "Details"}}]}
        restored = ProductProject.from_dict(project.to_dict())
        self.assertEqual(2, len(restored.variants))
        self.assertEqual("24寸", restored.variants[1].name)
        self.assertEqual(project.module_order, restored.module_order)
        self.assertIn("MAIN_DETAIL", restored.copy_versions)

    def test_white_main_image_has_no_copy_and_order_is_enforced(self) -> None:
        project = sample_project("outputs")
        briefs = BriefGenerator().generate(project)
        white = next(item for item in briefs if item.module_code == "MAIN_WHITE")
        self.assertTrue(all(not value for value in white.copy.values()))
        self.assertIn("纯白背景", white.design_brief)
        white_instance = next(item for item in project.module_instances if item.module_code == "MAIN_WHITE")
        detail_instance = next(item for item in project.module_instances if item.module_code == "MAIN_DETAIL")
        project.module_instances = [detail_instance, white_instance] + [
            item for item in project.module_instances if item.instance_id not in {detail_instance.instance_id, white_instance.instance_id}
        ]
        project.module_order = [item.instance_id for item in project.module_instances]
        self.assertTrue(any("第1张" in error for error in project.validate()))

    def test_custom_prompt_and_variant_summary_enter_brief(self) -> None:
        project = sample_project("outputs")
        detail = next(item for item in BriefGenerator().generate(project) if item.module_code == "MAIN_DETAIL")
        self.assertIn("macro close-up", detail.image_prompt)
        self.assertIn("20寸", detail.design_brief)
        self.assertIn("24寸", detail.design_brief)

    def test_repeated_module_instances_have_independent_prompts(self) -> None:
        project = sample_project("outputs")
        first = ModuleInstance("DETAIL_A", "MAIN_DETAIL", "主图", "轮组微距特写")
        second = ModuleInstance("DETAIL_B", "MAIN_DETAIL", "主图", "拉杆结构剖面")
        white = next(item for item in project.module_instances if item.module_code == "MAIN_WHITE")
        project.module_instances = [white, first, second]
        project.module_order = [item.instance_id for item in project.module_instances]
        project.selected_modules = {"主图": ["MAIN_WHITE", "MAIN_DETAIL", "MAIN_DETAIL"], "高级A+": [], "Brand Story": [], "品牌旗舰店": []}
        details = [item for item in BriefGenerator().generate(project) if item.module_code == "MAIN_DETAIL"]
        self.assertEqual(["DETAIL_A", "DETAIL_B"], [item.instance_id for item in details])
        self.assertIn("轮组微距", details[0].image_prompt)
        self.assertIn("拉杆结构", details[1].image_prompt)

    def test_provider_presets_and_separate_credentials(self) -> None:
        self.assertTrue({"openai", "cunai", "aliyun", "volcengine", "zhipu", "siliconflow", "baidu", "custom"}.issubset(TEXT_PROVIDER_PRESETS))
        self.assertTrue({"openai", "cunai", "aliyun", "volcengine", "zhipu", "siliconflow", "custom"}.issubset(IMAGE_PROVIDER_PRESETS))
        options = GenerationOptions(
            text_provider="custom",
            text_base_url="https://text.example/v1",
            text_api_key_env="NO_TEXT_KEY",
            image_provider="custom",
            image_base_url="https://image.example/v1",
            image_api_key_env="NO_IMAGE_KEY",
        )
        client = OpenAIClient(options=options, text_api_key="text-secret", image_api_key="image-secret")
        self.assertTrue(client.text_available)
        self.assertTrue(client.image_available)
        self.assertEqual(("url", "https://example.test/image.png"), client._find_image_result({"images": [{"url": "https://example.test/image.png"}]}))

    def test_cunai_file_configuration_and_python_chat_payload(self) -> None:
        preset = TEXT_PROVIDER_PRESETS["cunai"]
        self.assertEqual("https://www.cun.ai/v1", preset.base_url)
        self.assertEqual("/chat/completions", preset.endpoint)
        self.assertEqual("CUNAI_API_KEY", preset.api_key_env)
        self.assertEqual("claude-fable-5", preset.default_model)
        self.assertEqual("https://www.cun.ai/v1/models", model_catalog_url("cunai", "text", preset.base_url))
        raw = PROVIDER_CONFIG_PATH.read_text(encoding="utf-8")
        self.assertNotIn('"api_key":', raw)

        class StubCunClient(OpenAIClient):
            def _get_json(self, *args, **kwargs):
                return {"data": [{"id": "claude-fable-5"}, {"id": "model-b"}]}

            def _post_json(self, base_url, endpoint, payload, *args, **kwargs):
                self.request = (base_url, endpoint, payload)
                return {"choices": [{"message": {"content": "{\"ok\": true}"}}]}

        options = GenerationOptions(
            text_provider="cunai",
            text_protocol=preset.protocol,
            text_base_url=preset.base_url,
            text_endpoint=preset.endpoint,
            text_api_key_env=preset.api_key_env,
            text_model=preset.default_model,
            cache_enabled=False,
        )
        client = StubCunClient(options=options, text_api_key="test-key")
        self.assertEqual({"ok": True}, client.generate_json("Return JSON", preset.default_model))
        self.assertEqual("https://www.cun.ai/v1", client.request[0])
        self.assertEqual("/chat/completions", client.request[1])
        self.assertEqual("claude-fable-5", client.request[2]["model"])
        connection = client.test_connection("text")
        self.assertEqual(2, connection["model_count"])
        self.assertEqual("claude-fable-5", connection["tested_model"])
        stale = GenerationOptions(text_provider="cunai", text_base_url="https://stale.invalid", text_model="claude-fable-5")
        resolved = OpenAIClient.from_options(stale, text_api_key="test-key", build_fallbacks=False)
        self.assertEqual("https://www.cun.ai/v1", resolved.options.text_base_url)

    def test_cunai_image_provider_uses_shared_key_and_dynamic_models(self) -> None:
        preset = IMAGE_PROVIDER_PRESETS["cunai"]
        self.assertEqual("https://www.cun.ai/v1", preset.base_url)
        self.assertEqual("/images/generations", preset.endpoint)
        self.assertEqual("openai_images_url", preset.protocol)
        self.assertEqual("CUNAI_API_KEY", preset.api_key_env)
        self.assertEqual((), preset.models)
        self.assertEqual("https://www.cun.ai/v1/models", model_catalog_url("cunai", "image", preset.base_url))

    def test_variant_copy_preserves_attributes_and_marks_identity(self) -> None:
        from amazon_image_brief.gui import AmazonImageBriefApp

        source = ProductVariant("20寸", "CASE-20", "55×35×23cm", "Black", "2.8kg", "38L", "€89.99")
        copied = AmazonImageBriefApp._duplicate_variant_value(source)
        self.assertEqual("20寸（副本）", copied.name)
        self.assertEqual("CASE-20-COPY", copied.sku)
        self.assertEqual(source.size, copied.size)
        self.assertEqual(source.color, copied.color)
        self.assertEqual(source.weight, copied.weight)
        self.assertEqual(source.capacity, copied.capacity)
        self.assertEqual(source.price, copied.price)

    def test_custom_chat_protocol_parses_json_without_network(self) -> None:
        class StubClient(OpenAIClient):
            def _post_json(self, *args, **kwargs):
                return {"choices": [{"message": {"content": "```json\n{\"ok\": true}\n```"}}]}

        options = GenerationOptions(
            text_provider="custom",
            text_protocol="chat_completions",
            text_base_url="https://text.example/v1",
            text_endpoint="/chat/completions",
            text_api_key_env="NO_TEXT_KEY",
        )
        client = StubClient(options=options, text_api_key="temporary-secret")
        self.assertEqual({"ok": True}, client.generate_json("Return JSON", "custom-model"))

    def test_legacy_sample_migrates_duplicate_modules_to_unique_instances(self) -> None:
        sample_path = Path(__file__).resolve().parents[1] / "sample_suitcase_project.json"
        project = ProductProject.from_dict(json.loads(sample_path.read_text(encoding="utf-8")))
        repeated = [item for item in project.module_instances if item.module_code == "APLUS_SINGLE_TEXT"]
        self.assertEqual(2, len(repeated))
        self.assertEqual(2, len({item.instance_id for item in repeated}))


class ImportAndPaletteTests(unittest.TestCase):
    def test_market_template_export_and_import(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "template.xlsx"
            manager = MarketWorkbookManager()
            manager.export_template(path)
            workbook = load_workbook(path)
            workbook["竞品信息"].append(["B0123", "https://www.amazon.de/dp/B0123", "Brand", "€99", "Smooth wheels", "Clear layout", "Show capacity"])
            workbook["关键词词库"].append(["Koffer mit Rollen", "de", "category", 5, "core"])
            workbook.save(path)
            workbook.close()
            competitors, keywords = manager.import_data(path)
            self.assertEqual("B0123", competitors[0].asin)
            self.assertEqual("Koffer mit Rollen", keywords[0].term)
            self.assertEqual(5, keywords[0].priority)

    def test_logo_palette_ignores_white_background(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "logo.png"
            image = Image.new("RGB", (300, 160), "white")
            draw = ImageDraw.Draw(image)
            draw.rectangle((30, 30, 145, 130), fill="#123A63")
            draw.rectangle((155, 30, 270, 130), fill="#D88A2A")
            image.save(path)
            colors = extract_logo_palette(str(path), 4)
            self.assertGreaterEqual(len(colors), 2)
            self.assertFalse(any(color in {"#FFFFFF", "#FEFEFE"} for color in colors))


class V22RuntimeTests(unittest.TestCase):
    def test_ai_cache_round_trip_and_clear(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            cache = AICache(root, True)
            key = cache.key({"provider": "openai", "prompt": "safe"})
            cache.put_json(key, {"ok": True})
            self.assertEqual({"ok": True}, cache.get_json(key))
            source = root / "source.jpg"
            Image.new("RGB", (50, 50), "white").save(source)
            cache.put_image(key, source)
            copied = root / "copied.jpg"
            self.assertTrue(cache.copy_image_to(key, copied))
            self.assertTrue(copied.is_file())
            self.assertEqual(2, cache.clear())

    def test_text_request_automatically_fails_over(self) -> None:
        class FailingClient(OpenAIClient):
            def _generate_json_once(self, prompt: str, model: str):
                raise OpenAIError("primary unavailable")

        class WorkingClient(OpenAIClient):
            def _generate_json_once(self, prompt: str, model: str):
                return {"provider": "fallback"}

        options = GenerationOptions(auto_failover=True, cache_enabled=False)
        primary = FailingClient(options=options, text_api_key="primary")
        fallback_options = GenerationOptions(auto_failover=False, cache_enabled=False, text_model="fallback-model")
        primary.text_fallbacks = [WorkingClient(options=fallback_options, text_api_key="backup")]
        self.assertEqual({"provider": "fallback"}, primary.generate_json("Return JSON", options.text_model))

    def test_model_catalog_parser_and_cost_estimate(self) -> None:
        parsed = OpenAIClient._extract_model_ids({"data": [{"id": "b"}, {"id": "a"}, {"id": "a"}]})
        self.assertEqual(["a", "b"], parsed)
        project = sample_project("outputs")
        project.options.optimize_copy_with_ai = True
        project.options.generate_ai_images = True
        estimate = estimate_project_cost(project, module_count=2)
        self.assertEqual("USD", estimate.currency)
        self.assertEqual(2, estimate.image_count)
        self.assertGreater(estimate.total or 0, 0)

    def test_similarity_overflow_and_rule_preflight(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            image_path = Path(temporary) / "product.jpg"
            Image.new("RGB", (1200, 1200), "white").save(image_path)
            self.assertGreaterEqual(image_similarity(image_path, image_path), 99.0)
            project = sample_project("outputs", str(image_path))
            project.selected_modules = {"主图": ["MAIN_WHITE"], "高级A+": [], "Brand Story": [], "品牌旗舰店": []}
            project.module_instances = [next(item for item in project.module_instances if item.module_code == "MAIN_WHITE")]
            project.module_order = [project.module_instances[0].instance_id]
            briefs = BriefGenerator().generate(project)
            report = RuleLibrary.load().preflight(project, briefs)
            self.assertFalse(report.errors)
            self.assertEqual("DE", report.marketplace_code)
            long_brief = ModuleBrief("X", "高级A+", "APLUS_TEXT", "文本", "1464×600", "600×450", "", copy={"de": "Sehr lang " * 400})
            self.assertIn("de:", text_overflow_report(long_brief))

    def test_persistent_batch_queue_runs_and_restores(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            project_path = root / "project.json"
            project_path.write_text(json.dumps(sample_project(str(root / "out")).to_dict(), ensure_ascii=False), encoding="utf-8")
            queue_path = root / "queue.json"
            tasks = PersistentBatchQueue(queue_path)
            self.assertEqual(1, len(tasks.add_projects([str(project_path)])))

            def processor(project, progress, checkpoint):
                checkpoint()
                progress("完成测试", 1, 1)
                return {"output_dir": str(root / "result")}

            tasks.start(processor)
            deadline = time.time() + 3
            while tasks.is_running and time.time() < deadline:
                time.sleep(0.02)
            self.assertEqual("已完成", tasks.items[0].status)
            restored = PersistentBatchQueue(queue_path)
            self.assertEqual("已完成", restored.items[0].status)
            self.assertEqual(str(root / "result"), restored.items[0].result_dir)


class ArtifactTests(unittest.TestCase):
    def test_excel_export_contains_v2_columns_variants_and_versions(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            project = sample_project(str(root))
            briefs = BriefGenerator().generate(project)
            detail = next(item for item in briefs if item.module_code == "MAIN_DETAIL")
            detail.review_status = "已通过"
            detail.review_notes = "母语复核完成"
            detail.self_check_result = "结论：通过"
            project.copy_versions = {
                "MAIN_DETAIL": [
                    {"version": 1, "time": "2026-09-04 10:00:00", "reason": "初稿", "copy": dict(detail.copy)}
                ]
            }
            destination = root / "brief-v2.xlsx"
            ExcelExporter().export(project, briefs, destination)
            workbook = load_workbook(destination, read_only=False)
            self.assertEqual(
                ["项目参数", "主图", "高级A+", "Brand Story", "品牌旗舰店", "文案版本记录"],
                workbook.sheetnames,
            )
            headers = [workbook["主图"].cell(3, column).value for column in range(1, 29)]
            for required in (
                "模块实例ID", "用户指定Prompt", "文案版本", "审核状态", "ChatGPT自查结果/修改建议",
                "产品结构相似度", "OCR识别文字", "六语回译检查", "文案溢出检查", "图片文字合规检查", "Amazon规则预检",
            ):
                self.assertIn(required, headers)
            values = [workbook["项目参数"].cell(row, 1).value for row in range(1, workbook["项目参数"].max_row + 1)]
            self.assertIn("产品变体", values)

    def test_excel_includes_preflight_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            project = sample_project(temporary)
            briefs = BriefGenerator().generate(project)
            project.preflight_results = RuleLibrary.load().preflight(project, briefs).to_dict()
            destination = Path(temporary) / "with-preflight.xlsx"
            ExcelExporter().export(project, briefs, destination)
            workbook = load_workbook(destination, read_only=True)
            self.assertIn("Amazon规则预检", workbook.sheetnames)
            self.assertEqual("级别", workbook["Amazon规则预检"]["A2"].value)
            workbook.close()

    def test_image_composition_outputs_jpeg(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            base = root / "base.jpg"
            composite = root / "composite.jpg"
            create_placeholder(base, "尺寸与容量", "test prompt", "square")
            create_german_composite(base, composite, "Mehr Platz für jede Reise")
            with Image.open(composite) as image:
                self.assertEqual("JPEG", image.format)
                self.assertGreaterEqual(image.width, 1000)

    def test_service_dry_run_creates_excel_and_reference_copy(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "suitcase.png"
            Image.new("RGB", (500, 500), "white").save(source)
            project = sample_project(str(root / "out"), str(source))
            project.selected_modules = {"主图": ["MAIN_WHITE"], "高级A+": [], "Brand Story": [], "品牌旗舰店": []}
            project.module_instances = [next(item for item in project.module_instances if item.module_code == "MAIN_WHITE")]
            project.module_order = [project.module_instances[0].instance_id]
            result = GenerationService(root).run(project)
            self.assertTrue(Path(result["excel"]).is_file())
            self.assertTrue((Path(result["output_dir"]) / "preflight_report.json").is_file())
            self.assertTrue((Path(result["output_dir"]) / "qa_report.json").is_file())
            self.assertEqual(1, len(list((Path(result["output_dir"]) / "reference_images").glob("*.jpg"))))
            workbook = load_workbook(result["excel"], read_only=False)
            self.assertEqual(1, len(workbook["主图"]._images))
            self.assertIn("Amazon规则预检", workbook.sheetnames)
            workbook.close()

    def test_copy_regeneration_and_self_check(self) -> None:
        project = sample_project("outputs")
        generator = BriefGenerator(FakeAIClient())
        briefs = generator.generate(project)
        detail = next(item for item in briefs if item.module_code == "MAIN_DETAIL")
        rewritten = generator.regenerate_copy(project, detail, "缩短标题")
        self.assertEqual("新版-de", rewritten["copy"]["de"])
        checks = generator.self_check(project, [detail])
        key = detail.instance_id or detail.module_code
        self.assertIn("建议修改", checks[key])
        self.assertIn("缩短支持句", checks[key])

    def test_partial_image_regeneration_refreshes_excel(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "suitcase.png"
            Image.new("RGB", (500, 500), "white").save(source)
            project = sample_project(str(root / "out"), str(source))
            project.selected_modules = {"主图": ["MAIN_WHITE", "MAIN_DETAIL"], "高级A+": [], "Brand Story": [], "品牌旗舰店": []}
            project.module_instances = [
                next(item for item in project.module_instances if item.module_code == "MAIN_WHITE"),
                next(item for item in project.module_instances if item.module_code == "MAIN_DETAIL"),
            ]
            project.module_order = [item.instance_id for item in project.module_instances]
            briefs = BriefGenerator().generate(project)
            service = GenerationService(root, FakeAIClient())
            result = service.run(project, briefs=briefs)
            detail = next(item for item in briefs if item.module_code == "MAIN_DETAIL")
            refreshed = service.regenerate_image(project, briefs, detail.instance_id, Path(result["output_dir"]))
            self.assertTrue(Path(detail.ai_effect_image).is_file())
            self.assertIn("regen", Path(detail.ai_effect_image).name)
            self.assertTrue(Path(refreshed["excel"]).is_file())

    def test_mask_regeneration_passes_mask_and_structure_lock(self) -> None:
        class MaskAwareClient:
            available = True
            image_available = True
            text_available = False

            def __init__(self) -> None:
                self.mask_path = None
                self.prompt = ""

            def generate_image(self, prompt, destination, model, quality, ratio, reference_images, mask_path=None, use_cache=True):
                self.mask_path = mask_path
                self.prompt = prompt
                self.use_cache = use_cache
                create_placeholder(Path(destination), "蒙版重生", prompt, ratio)

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "suitcase.png"
            mask = root / "mask.png"
            Image.new("RGB", (1200, 1200), "white").save(source)
            Image.new("RGBA", (1200, 1200), (0, 0, 0, 255)).save(mask)
            project = sample_project(str(root / "out"), str(source))
            project.selected_modules = {"主图": ["MAIN_WHITE", "MAIN_DETAIL"], "高级A+": [], "Brand Story": [], "品牌旗舰店": []}
            project.module_instances = [
                next(item for item in project.module_instances if item.module_code == "MAIN_WHITE"),
                next(item for item in project.module_instances if item.module_code == "MAIN_DETAIL"),
            ]
            project.module_order = [item.instance_id for item in project.module_instances]
            briefs = BriefGenerator().generate(project)
            detail = next(item for item in briefs if item.module_code == "MAIN_DETAIL")
            client = MaskAwareClient()
            result = GenerationService(root, client).regenerate_image(project, briefs, detail.instance_id, root / "existing", mask)
            self.assertEqual(mask, client.mask_path)
            self.assertFalse(client.use_cache)
            self.assertIn("PRODUCT STRUCTURE LOCK", client.prompt)
            self.assertTrue(Path(result["excel"]).is_file())


if __name__ == "__main__":
    unittest.main()
