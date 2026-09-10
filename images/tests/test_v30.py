from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
import json
import tempfile
import tkinter as tk
import unittest
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook

from amazon_image_brief.ai_client import OpenAIClient, OpenAIError
from amazon_image_brief.models import GenerationOptions, ModuleInstance, ProductProject
from amazon_image_brief.product_context import ProductContext, context_fingerprint, product_facts
from amazon_image_brief.creative_ai import CreativeAI, validate_ai_plan, recipe_for
from amazon_image_brief.creative_planner import plan_modules
from amazon_image_brief.brief_generator import BriefGenerator
from amazon_image_brief.service import GenerationService
from amazon_image_brief.navigation import expand_navigation, resize_navigation
from amazon_image_brief.reference_inputs import reference_inputs
from amazon_image_brief.gui import AmazonImageBriefApp
from amazon_image_brief.profile_store import ProfileStore
from test_v28 import PlannedAI, project_at


class ContextAI(PlannedAI):
    image_available = True

    def __init__(self):
        super().__init__()
        self.calls, self.image_calls = [], []

    def analyze_images_json(self, prompt, paths, model, **kwargs):
        self.calls.append(('vision', prompt, [str(p) for p in paths]))
        if 'STYLE REFERENCES' in prompt:
            return {'layout_description': 'Centered headline above actual product', 'visual_style': 'editorial'}
        return {'summary_zh': '已识别产品实物：外观以图片为准', 'unknowns': ['材质等级不能从照片推断']}

    def generate_json(self, prompt, model):
        self.calls.append(('text', prompt))
        if 'summary_zh (string)' in prompt:
            return {'summary_zh': '已理解录入资料，无实拍图片'}
        if 'Return JSON {"prompt"' in prompt:
            return {'prompt': '优化后的本帧专属提示词：保留可验证的产品事实', 'rationale_zh': '结合本帧主卖点'}
        if 'Return JSON with rationale_zh' in prompt:
            plan = json.loads(prompt.split('\nCurrent module plan: ')[1].split('\nOriginal per-image prompts: ')[0])
            return {'rationale_zh': '顶端标题与产品部位说明', 'visual': '产品近景', **{d: plan[d] for d in ('pc', 'mobile')}}
        return super().generate_json(prompt, model)

    def generate_image(self, prompt, destination, **kwargs):
        self.calls.append(('image', prompt))
        self.image_calls.append({'prompt': prompt, **kwargs})
        destination.parent.mkdir(parents=True, exist_ok=True)
        Image.new('RGB', (1200, 1200), 'white').save(destination, 'JPEG')
        return destination


def photos(root, count=2):
    result = []
    for n in range(count):
        path = root/f'product-{n}.png'
        Image.new('RGB', (800, 1000), ['red', 'blue', 'green'][n % 3]).save(path)
        result.append(str(path))
    return result


class ContextAndGenerationTests(unittest.TestCase):
    def test_context_all_photos_batches_cache_and_invalidation(self):
        with tempfile.TemporaryDirectory() as temp:
            project = project_at(temp)
            project.options.context_before_generation = True
            project.product_image_paths = photos(Path(temp), 5)
            project.price = 'EUR 79.99'
            client, events = ContextAI(), []
            worker = ProductContext(client, events.append)
            worker.prepare(project)
            self.assertEqual([1, 1, 1, 1, 1], [len(c[2]) for c in client.calls])
            self.assertEqual(project.product_image_paths, [p for c in client.calls for p in c[2]])
            self.assertIn('79.99', client.calls[0][1])
            worker.prepare(project)
            self.assertEqual(5, len(client.calls))
            project.price = 'EUR 89.99'
            worker.prepare(project)
            self.assertEqual(10, len(client.calls))
            self.assertEqual('complete', events[-1]['context']['status'])
            old = context_fingerprint(project)
            Image.new('RGB', (800, 1000), 'yellow').save(project.product_image_paths[0])
            self.assertNotEqual(old, context_fingerprint(project))
            project.options.text_api_key_env = 'NEVER_INCLUDE_SECRET_ENV'
            self.assertNotIn('NEVER_INCLUDE', json.dumps(product_facts(project)))

    def test_context_failure_stops_generation_before_image_request(self):
        with tempfile.TemporaryDirectory() as temp:
            project = project_at(Path(temp)/'out')
            project.product_image_paths = photos(Path(temp), 1)
            project.options.context_before_generation = project.options.generate_ai_images = True
            client = ContextAI()
            client.analyze_images_json = lambda *args, **kwargs: {'invalid': True}
            with self.assertRaises(OpenAIError):
                GenerationService(Path(temp), client).run(project)
            self.assertEqual([], client.image_calls)

    def test_copy_context_precedes_requests_and_explicit_user_prompt_survives(self):
        with tempfile.TemporaryDirectory() as temp:
            project = project_at(temp)
            project.options.context_before_generation = project.options.optimize_copy_with_ai = True
            project.product_image_paths = photos(Path(temp))
            project.module_recipes['detail1'] = {'copy_prompt': '只突出真实轮组，标题不要重复'}
            client = ContextAI()
            briefs = BriefGenerator(client).generate(project)
            self.assertEqual('vision', client.calls[0][0])
            self.assertEqual(2, sum(c[0] == 'vision' for c in client.calls))
            self.assertIn('只突出真实轮组', briefs[1].effective_copy_prompt)
            self.assertIn('已识别产品实物', briefs[2].effective_copy_prompt)
            self.assertNotEqual(briefs[1].copy['de'], briefs[2].copy['de'])

    def test_image_requests_use_approved_copy_user_prompt_and_style_then_export_trace(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            project = project_at(root/'out')
            project.options.generate_ai_images = project.options.context_before_generation = True
            project.product_image_paths = photos(root, 2)
            project.module_recipes['detail1'] = {'image_prompt': 'USER GOLD WHEEL CLOSEUP', 'style_reference_paths': [project.product_image_paths[1]], 'render_text': 'native'}
            briefs = BriefGenerator(ContextAI()).generate(project)
            briefs[1].copy['de'] = 'MEIN EXAKTER TEXT\nKein Ersatz'
            briefs[1].image_prompt = 'STALE PROMPT MUST NOT WIN'
            client, events = ContextAI(), []
            result = GenerationService(root, client, on_image_progress=events.append).run(project, briefs=briefs)
            prompt = client.image_calls[1]['prompt']
            self.assertIn('USER GOLD WHEEL CLOSEUP', prompt)
            self.assertIn('MEIN EXAKTER TEXT\nKein Ersatz', prompt)
            self.assertNotIn('STALE PROMPT', prompt)
            self.assertIn('Centered headline', prompt)
            self.assertIn('Render the EXACT', prompt)
            self.assertIn('no text', client.image_calls[0]['prompt'])
            self.assertEqual(prompt, result['briefs'][1].effective_image_prompt)
            self.assertTrue(any(e['stage'] == 'complete' for e in events))
            workbook = load_workbook(result['excel'])
            self.assertIn('V3提示词与上下文', workbook.sheetnames)
            workbook.close()

    def test_multiple_actual_images_in_both_wire_protocols(self):
        with tempfile.TemporaryDirectory() as temp:
            paths = [Path(p) for p in photos(Path(temp), 3)]
            for protocol in ('chat_completions', 'responses'):
                client = OpenAIClient(options=GenerationOptions(text_protocol=protocol, cache_enabled=False, api_debug_enabled=False), text_api_key='test-only')
                reply = {'choices': [{'message': {'content': '{"ok":true}'}}]} if protocol == 'chat_completions' else {'output': [{'type': 'message', 'content': [{'type': 'output_text', 'text': '{"ok":true}'}]}]}
                with patch.object(client, '_post_json', return_value=reply) as post:
                    self.assertEqual({'ok': True}, client.analyze_images_json('inspect all', paths))
                payload = post.call_args.args[2]
                parts = payload['messages' if protocol == 'chat_completions' else 'input'][0]['content']
                self.assertEqual(4, len(parts))
                self.assertTrue(all('data:image/jpeg;base64,' in json.dumps(part) for part in parts[1:]))

    def test_reference_board_keeps_all_roles_and_native_multi_image_route(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            paths = photos(root, 3)
            refs, manifest = reference_inputs(paths[:2], paths[2:], ContextAI(), root/'cache')
            self.assertEqual(3, len(manifest['sources']))
            self.assertEqual('STYLE ONLY', manifest['sources'][-1]['role'])
            self.assertTrue(Path(refs[0]).is_file())
            client = SimpleNamespace(options=GenerationOptions(image_provider='openai', image_protocol='openai_images'))
            refs, manifest = reference_inputs(paths[:2], paths[2:], client, root/'cache')
            self.assertEqual(paths, refs)

    def test_force_context_refresh_bypasses_response_cache_and_restores_setting(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            project = project_at(temp)
            project.options.context_before_generation = True
            project.product_image_paths = photos(root, 1)
            client = OpenAIClient(options=GenerationOptions(cache_enabled=True, api_debug_enabled=False, text_protocol='chat_completions'), text_api_key='test-key', cache_dir=root/'cache')
            response = {'choices': [{'message': {'content': '{"summary_zh":"verified"}'}}]}
            with patch.object(client, '_post_json', return_value=response) as post:
                worker = ProductContext(client)
                worker.prepare(project)
                worker.prepare(project)
                self.assertEqual(1, post.call_count)
                worker.prepare(project, force=True)
                self.assertEqual(2, post.call_count)
            self.assertTrue(client.cache.enabled)

    def test_navigation_generates_one_jpg_per_frame_and_regenerates_only_selected(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            project = project_at(root/'out')
            project.module_instances = [ModuleInstance('nav', 'APLUS_NAV_CAROUSEL', '高级A+')]
            project.product_image_paths = photos(root, 1)
            project.options.generate_ai_images = True
            for frame in project.normalized_module_instances():
                project.module_recipes[frame.instance_id] = {'image_prompt': f'INDEPENDENT FRAME {frame.frame_index}'}
            client = ContextAI()
            service = GenerationService(root, client)
            result = service.run(project)
            self.assertEqual(4, len(client.image_calls))
            self.assertEqual(4, len({b.ai_effect_image for b in result['briefs']}))
            for n, call in enumerate(client.image_calls, 1):
                self.assertIn(f'INDEPENDENT FRAME {n}', call['prompt'])
                self.assertTrue(Path(result['briefs'][n-1].ai_effect_image).is_file())
            before = [b.ai_effect_image for b in result['briefs']]
            service.regenerate_image(project, result['briefs'], result['briefs'][2].instance_id, Path(result['output_dir']))
            self.assertEqual(5, len(client.image_calls))
            self.assertNotEqual(before[2], result['briefs'][2].ai_effect_image)
            self.assertEqual(before[:2], [b.ai_effect_image for b in result['briefs'][:2]])
            self.assertEqual({}, client.debug_context)


class RecipesAndNavigationTests(unittest.TestCase):
    def test_nav_default_four_stable_custom_count_and_round_trip(self):
        project = project_at('unused')
        project.module_instances = [ModuleInstance('nav', 'APLUS_NAV_CAROUSEL', '高级A+')]
        frames = project.normalized_module_instances()
        self.assertEqual(4, len(frames))
        self.assertEqual('nav', frames[0].instance_id)
        self.assertEqual([1, 2, 3, 4], [i.frame_index for i in frames])
        grown = resize_navigation(frames, 'nav', 6)
        self.assertEqual([i.instance_id for i in frames], [i.instance_id for i in grown[:4]])
        project.module_instances = grown
        project.module_recipes = {i.instance_id: {'image_prompt': f'FRAME {i.frame_index}'} for i in grown}
        restored = ProductProject.from_dict(project.to_dict())
        briefs = BriefGenerator(ContextAI()).generate(restored)
        self.assertEqual(6, len(briefs))
        for brief in briefs:
            self.assertIn(f'FRAME {brief.frame_index}', GenerationService._generation_prompt(restored, brief))
        one = resize_navigation(grown, 'nav', 1)
        self.assertEqual(1, len(expand_navigation(one)))
        with self.assertRaises(ValueError):
            resize_navigation(grown, 'nav', 0)

    def test_prompt_optimizer_updates_original_field_and_history_per_frame(self):
        project = project_at('unused')
        project.module_recipes['detail1'] = {'copy_prompt': '原有指定提示词'}
        events = []
        CreativeAI(ContextAI(), on_item=events.append).run(project, ['detail1', 'detail2'], 'copy')
        self.assertEqual(2, len(events))
        self.assertIn('优化后', project.module_recipes['detail1']['copy_prompt'])
        self.assertEqual('原有指定提示词', project.module_recipes['detail1']['prompt_history'][0]['previous'])
        self.assertNotIn('image_prompt', project.module_recipes.get('white', {}))

    def test_invalid_prompt_preserves_original_and_partial_success(self):
        project = project_at('unused')
        project.module_recipes['detail2'] = {'image_prompt': 'KEEP'}
        client = ContextAI()
        original = client.generate_json
        def response(prompt, model):
            if 'KEEP' in prompt:
                return {'prompt': ''}
            return original(prompt, model)
        client.generate_json = response
        with self.assertRaises(OpenAIError):
            CreativeAI(client).run(project, ['detail1', 'detail2'], 'image')
        self.assertIn('优化后', project.module_recipes['detail1']['image_prompt'])
        self.assertEqual('KEEP', project.module_recipes['detail2']['image_prompt'])

    def test_ai_layout_validation_and_saved_ai_layout_survives_poster_sync(self):
        project = project_at('unused')
        project.module_instances = [ModuleInstance('a', 'APLUS_FULL_IMAGE', '高级A+')]
        base = plan_modules(project)['a']
        answer = {device: deepcopy(base[device]) for device in ('pc', 'mobile')}
        answer['rationale_zh'] = '按内容确定标题与特写说明的位置'
        plan = validate_ai_plan(answer, base)
        project.creative_plans['a'] = plan
        project.options.aplus_continuous = True
        restored = plan_modules(ProductProject.from_dict(project.to_dict()))['a']
        self.assertEqual('ai', restored['layout'])
        self.assertEqual(plan['pc']['elements'], restored['pc']['elements'])
        answer['pc']['elements'][1]['box'] = answer['pc']['elements'][0]['box']
        with self.assertRaises(OpenAIError):
            validate_ai_plan(answer, base)

    def test_profile_manages_style_assets_and_preserves_prompts(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            project = project_at(temp)
            project.module_recipes['detail1'] = {'image_prompt': 'KEEP RECIPE', 'style_reference_paths': photos(root, 1)}
            store = ProfileStore(root/'user_data')
            # Public save/load contract also covers deeply nested recipe image lists.
            record = {'project': project.to_dict(), 'briefs': [], 'ui': {}}
            store.save(record)
            restored = store.load()
            recipe = restored['project']['module_recipes']['detail1']
            self.assertEqual('KEEP RECIPE', recipe['image_prompt'])
            self.assertTrue(Path(recipe['style_reference_paths'][0]).is_file())


class V30DesktopTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = tk.Tk()
        self.root.withdraw()
        self.app = AmazonImageBriefApp(self.root, Path(self.temp.name))
        self.app.populate_project(project_at(self.temp.name))

    def tearDown(self):
        for control in self.app.generation_controls.values():
            control.close()
        for callback in self.root.tk.call('after', 'info'):
            self.root.after_cancel(callback)
        self.root.destroy()
        self.temp.cleanup()

    def test_defaults_language_comparison_and_recipe_autosave(self):
        app = self.app
        app.current_briefs = BriefGenerator(ContextAI()).generate(project_at(self.temp.name))
        app.current_review_code = 'detail1'
        app._load_review_editor()
        self.assertEqual('de', app.current_review_language)
        self.assertEqual('zh — 中文', app.compare_language_var.get())
        self.assertIn(app.current_briefs[1].copy['zh'], app.translation_compare_text.get('1.0', 'end-1c'))
        self.assertIn('暂无逐语回译', app.translation_compare_text.get('1.0', 'end-1c'))
        panel = app.recipe_panels['image']
        panel['var'].set(next(k for k, v in panel['choices'].items() if v == 'detail1'))
        app._select_recipe('image')
        panel['editor'].insert('1.0', '指定镜头与布局')
        project = app.collect_project()
        self.assertEqual('指定镜头与布局', project.module_recipes['detail1']['image_prompt'])
        app.populate_project(project)
        self.assertEqual('指定镜头与布局', app.module_recipes['detail1']['image_prompt'])

    def test_optimized_prompt_writes_same_visible_field_and_locks_during_work(self):
        app = self.app
        panel = app.recipe_panels['copy']
        iid = panel['id']
        app.busy = True
        app._start_controls('v3:copy')
        app._receive_recipe({'id': iid, 'action': 'copy', 'recipe': {'copy_prompt': 'AI优化结果'}, 'done': 1, 'total': 1})
        self.assertEqual('AI优化结果', panel['editor'].get('1.0', 'end-1c'))
        self.assertEqual('disabled', str(panel['editor']['state']))
        app.busy = False
        app._finish_controls()
        self.assertEqual('normal', str(panel['editor']['state']))

    def test_add_and_resize_nav_from_ui_with_independent_frames(self):
        from amazon_image_brief.catalogs import CATALOGS
        app = self.app
        listing = app.module_lists['高级A+']
        index = next(n for n, spec in enumerate(CATALOGS['高级A+']) if spec.code == 'APLUS_NAV_CAROUSEL')
        listing.selection_clear(0, 'end')
        listing.selection_set(index)
        app.add_module_instance('高级A+')
        frames = [i for i in app.module_instances if i.parent_id]
        self.assertEqual(4, len(frames))
        app.nav_count_var.set('6')
        app.module_selected_tree.selection_set(frames[0].instance_id)
        app.resize_selected_navigation()
        self.assertEqual(6, len([i for i in app.module_instances if i.parent_id]))

    def test_legacy_confirm_prompt_updates_image_recipe_without_overwriting_copy_override(self):
        app = self.app
        app.module_recipes['detail1'] = {'copy_prompt': '独立文案要求', 'image_prompt': '旧图片要求'}
        app.confirm_tree.selection_set('detail1')
        app._on_confirm_select()
        app.module_prompt_text.delete('1.0', 'end')
        app.module_prompt_text.insert('1.0', '新构图要求')
        app._save_module_prompt()
        project = app.collect_project()
        self.assertEqual('新构图要求', project.module_recipes['detail1']['image_prompt'])
        self.assertEqual('独立文案要求', project.module_recipes['detail1']['copy_prompt'])


if __name__ == '__main__':
    unittest.main()
